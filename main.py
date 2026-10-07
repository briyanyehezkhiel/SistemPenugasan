import tkinter as tk
import random
import datetime
from tkinter import ttk, messagebox, simpledialog, filedialog
from datetime import datetime
import calendar
import sqlite3
import math
import time
import re

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None
    load_workbook = None
    Font = PatternFill = Alignment = Border = Side = None
    get_column_letter = None

from database import get_connection, create_database


# =========================================================
# KONFIGURASI
# =========================================================

BAGIAN = []

NAMA_BULAN = [
    "",
    "Januari",
    "Februari",
    "Maret",
    "April",
    "Mei",
    "Juni",
    "Juli",
    "Agustus",
    "September",
    "Oktober",
    "November",
    "Desember"
]


# =========================================================
# HELPER
# =========================================================

def rupiah(nilai):

    return "Rp {:,.0f}".format(
        nilai
    ).replace(",", ".")


def parse_rupiah(teks):

    teks = (
        teks
        .replace("Rp", "")
        .replace(".", "")
        .replace(",", "")
        .strip()
    )

    if not teks:
        return 0

    return int(teks)


def nilai_excel_teks(nilai):
    """Mengubah nilai Excel menjadi teks tanpa mengubah data yang sudah berupa teks."""
    if nilai is None:
        return ""
    if isinstance(nilai, bool):
        return str(nilai)
    if isinstance(nilai, int):
        return str(nilai)
    if isinstance(nilai, float):
        if nilai.is_integer():
            return str(int(nilai))
        return str(nilai).strip()
    return str(nilai).strip()


def parse_tarif_excel(nilai):
    """Membaca tarif dari kolom tarif Excel (mis. 80000 atau Rp 80.000)."""
    teks = nilai_excel_teks(nilai)
    if not teks:
        return 0
    try:
        return parse_rupiah(teks)
    except (ValueError, TypeError):
        # Ambil digit saja untuk format Excel yang tidak standar.
        digit = re.sub(r"[^0-9]", "", teks)
        return int(digit) if digit else 0


def golongan_utama_dari(teks):
    """Mengambil golongan utama dari golongan asli, mis. 'III d' -> 'III'."""
    teks = (teks or "").strip().upper()
    if not teks:
        return ""
    return teks.split()[0]


def kode_bagian_berikutnya(kode_list):
    """Menghasilkan kode pertama yang belum dipakai: A, B, C, ..."""
    dipakai = {str(k).strip().upper() for k in kode_list if k}
    nomor = 0
    while True:
        n = nomor
        hasil = ""
        while True:
            hasil = chr(ord("A") + (n % 26)) + hasil
            n = n // 26 - 1
            if n < 0:
                break
        if hasil not in dipakai:
            return hasil
        nomor += 1


def ensure_bagian_schema():
    """Migrasi ringan untuk menambahkan PJ wajib tanpa menghapus data lama."""
    conn = get_connection()
    cursor = conn.cursor()
    columns = [row[1] for row in cursor.execute("PRAGMA table_info(bagian)").fetchall()]
    if "penanggung_jawab_id" not in columns:
        cursor.execute("ALTER TABLE bagian ADD COLUMN penanggung_jawab_id INTEGER")
    if "nama_bagian" not in columns:
        cursor.execute("ALTER TABLE bagian ADD COLUMN nama_bagian TEXT")

    # Hari libur disimpan per periode (bulan/tahun), sehingga setiap bulan
    # dapat memiliki tanggal merah yang berbeda. Hari Minggu tetap nonaktif
    # otomatis dan tidak perlu disimpan di tabel ini.
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS hari_libur (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            periode_id INTEGER NOT NULL,
            hari INTEGER NOT NULL,
            UNIQUE(periode_id, hari)
        )
    """)

    conn.commit()
    conn.close()


# =========================================================
# DATABASE
# =========================================================

def get_or_create_periode(bulan, tahun):

    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute("""
        INSERT OR IGNORE INTO periode
        (bulan, tahun)
        VALUES (?, ?)
    """, (bulan, tahun))

    periode = cursor.execute("""
        SELECT id FROM periode
        WHERE bulan = ? AND tahun = ?
    """, (bulan, tahun)).fetchone()

    periode_id = periode[0]

    # Periode baru minimal memiliki satu kegiatan.
    # Kegiatan berikutnya dibuat lewat tombol + Tambah Kegiatan.
    jumlah = cursor.execute(
        "SELECT COUNT(*) FROM bagian WHERE periode_id = ?",
        (periode_id,)
    ).fetchone()[0]

    if jumlah == 0:
        cursor.execute("""
            INSERT INTO bagian
            (periode_id, kode, kegiatan, lokasi, dana, penanggung_jawab_id)
            VALUES (?, 'A', '', '', 0, NULL)
        """, (periode_id,))

    conn.commit()
    conn.close()

    return periode_id


def load_bagian_codes(periode_id):
    conn = get_connection()
    cursor = conn.cursor()
    rows = cursor.execute(
        "SELECT kode FROM bagian WHERE periode_id = ? ORDER BY id",
        (periode_id,)
    ).fetchall()
    conn.close()
    return [row[0] for row in rows]


def load_hari_libur(periode_id):
    """Mengambil tanggal libur manual untuk periode aktif."""
    conn = get_connection()
    rows = conn.execute(
        "SELECT hari FROM hari_libur WHERE periode_id = ? ORDER BY hari",
        (periode_id,)
    ).fetchall()
    conn.close()
    return {row[0] for row in rows}


def simpan_hari_libur(periode_id, hari_libur):
    """Menyimpan ulang daftar hari libur manual untuk satu periode."""
    conn = get_connection()
    try:
        conn.execute(
            "DELETE FROM hari_libur WHERE periode_id = ?",
            (periode_id,)
        )
        for hari in sorted(set(hari_libur)):
            conn.execute(
                "INSERT INTO hari_libur (periode_id, hari) VALUES (?, ?)",
                (periode_id, hari)
            )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def nama_hari_libur(periode_id):
    hari = load_hari_libur(periode_id)
    if not hari:
        return "Tidak ada"
    return ", ".join(str(x) for x in sorted(hari))


def load_pegawai():

    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            id,
            nama,
            nip,
            pangkat,
            golongan_asli,
            golongan_utama,
            jabatan,
            tarif
        FROM pegawai
        WHERE aktif = 1
        ORDER BY nama
    """)

    data = cursor.fetchall()

    conn.close()

    return data


def load_bagian(periode_id, kode):

    conn = get_connection()
    cursor = conn.cursor()

    data = cursor.execute("""
        SELECT
            id,
            kode,
            kegiatan,
            lokasi,
            dana,
            penanggung_jawab_id
        FROM bagian
        WHERE periode_id = ?
        AND kode = ?
    """, (
        periode_id,
        kode
    )).fetchone()

    conn.close()

    return data


def load_anggota(bagian_id):
    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT
            p.id,
            p.nama
        FROM anggota_bagian ab
        JOIN pegawai p
            ON p.id = ab.pegawai_id
        WHERE ab.bagian_id = ?
        ORDER BY p.nama
        """,
        (bagian_id,)
    )

    data = cursor.fetchall()

    conn.close()

    return data
def load_riwayat_turun(periode_id):
    """Mengambil jumlah penugasan pegawai pada BULAN/periode aktif."""
    conn = get_connection()
    try:
        rows = conn.execute(
            """
            SELECT pegawai_id, COUNT(*) AS bulan_ini
            FROM jadwal
            WHERE periode_id = ?
            GROUP BY pegawai_id
            """,
            (periode_id,)
        ).fetchall()
        return {
            row[0]: {"bulan_ini": row[1] or 0}
            for row in rows
        }
    finally:
        conn.close()


def get_tahun_penugasan():
    """Daftar tahun yang mempunyai data periode, terbaru lebih dahulu."""
    conn = get_connection()
    try:
        rows = conn.execute(
            "SELECT DISTINCT tahun FROM periode ORDER BY tahun DESC"
        ).fetchall()
        return [int(row[0]) for row in rows]
    finally:
        conn.close()


def get_ringkasan_periode(periode_id):
    """Ringkasan data penugasan untuk satu bulan/periode."""
    conn = get_connection()
    try:
        kegiatan = conn.execute("SELECT COUNT(*) FROM bagian WHERE periode_id = ?", (periode_id,)).fetchone()[0]
        anggota = conn.execute(
            "SELECT COUNT(*) FROM anggota_bagian WHERE bagian_id IN (SELECT id FROM bagian WHERE periode_id = ?)",
            (periode_id,)
        ).fetchone()[0]
        jadwal = conn.execute("SELECT COUNT(*) FROM jadwal WHERE periode_id = ?", (periode_id,)).fetchone()[0]
        row = conn.execute("SELECT bulan, tahun FROM periode WHERE id = ?", (periode_id,)).fetchone()
        return {
            "periode": 1 if row else 0,
            "bulan": row[0] if row else None,
            "tahun": row[1] if row else None,
            "kegiatan": kegiatan,
            "anggota": anggota,
            "jadwal": jadwal,
        }
    finally:
        conn.close()


def hapus_data_penugasan_periode(periode_id):
    """Menghapus seluruh data penugasan satu bulan, tanpa menghapus master pegawai."""
    conn = get_connection()
    try:
        bagian_ids = [r[0] for r in conn.execute(
            "SELECT id FROM bagian WHERE periode_id = ?", (periode_id,)
        ).fetchall()]
        if bagian_ids:
            ph = ",".join("?" * len(bagian_ids))
            conn.execute(f"DELETE FROM jadwal WHERE bagian_id IN ({ph})", bagian_ids)
            conn.execute(f"DELETE FROM anggota_bagian WHERE bagian_id IN ({ph})", bagian_ids)
        conn.execute("DELETE FROM jadwal WHERE periode_id = ?", (periode_id,))
        conn.execute("DELETE FROM hari_libur WHERE periode_id = ?", (periode_id,))
        conn.execute("DELETE FROM bagian WHERE periode_id = ?", (periode_id,))
        conn.execute("DELETE FROM periode WHERE id = ?", (periode_id,))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def get_ringkasan_tahun(tahun):
    """Ringkasan data penugasan per tahun untuk jendela pembersihan."""
    conn = get_connection()
    try:
        periode_ids = [
            r[0] for r in conn.execute(
                "SELECT id FROM periode WHERE tahun = ?", (tahun,)
            ).fetchall()
        ]
        if not periode_ids:
            return {"periode": 0, "kegiatan": 0, "anggota": 0, "jadwal": 0}
        placeholders = ",".join("?" * len(periode_ids))
        kegiatan = conn.execute(
            f"SELECT COUNT(*) FROM bagian WHERE periode_id IN ({placeholders})",
            periode_ids
        ).fetchone()[0]
        anggota = conn.execute(
            f"SELECT COUNT(*) FROM anggota_bagian WHERE bagian_id IN (SELECT id FROM bagian WHERE periode_id IN ({placeholders}))",
            periode_ids
        ).fetchone()[0]
        jadwal = conn.execute(
            f"SELECT COUNT(*) FROM jadwal WHERE periode_id IN ({placeholders})",
            periode_ids
        ).fetchone()[0]
        return {"periode": len(periode_ids), "kegiatan": kegiatan, "anggota": anggota, "jadwal": jadwal}
    finally:
        conn.close()


def hapus_data_penugasan_tahun(tahun):
    """Menghapus seluruh data penugasan satu tahun, tanpa menghapus master pegawai."""
    conn = get_connection()
    try:
        periode_ids = [
            r[0] for r in conn.execute(
                "SELECT id FROM periode WHERE tahun = ?", (tahun,)
            ).fetchall()
        ]
        if not periode_ids:
            return 0
        placeholders = ",".join("?" * len(periode_ids))
        bagian_ids = [
            r[0] for r in conn.execute(
                f"SELECT id FROM bagian WHERE periode_id IN ({placeholders})",
                periode_ids
            ).fetchall()
        ]
        if bagian_ids:
            ph_b = ",".join("?" * len(bagian_ids))
            conn.execute(f"DELETE FROM jadwal WHERE bagian_id IN ({ph_b})", bagian_ids)
            conn.execute(f"DELETE FROM anggota_bagian WHERE bagian_id IN ({ph_b})", bagian_ids)
        conn.execute(f"DELETE FROM jadwal WHERE periode_id IN ({placeholders})", periode_ids)
        conn.execute(f"DELETE FROM hari_libur WHERE periode_id IN ({placeholders})", periode_ids)
        conn.execute(f"DELETE FROM bagian WHERE periode_id IN ({placeholders})", periode_ids)
        conn.execute(f"DELETE FROM periode WHERE id IN ({placeholders})", periode_ids)
        conn.commit()
        return len(periode_ids)
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def hapus_semua_data_penugasan():
    """Menghapus seluruh histori penugasan semua tahun, master pegawai tetap."""
    conn = get_connection()
    try:
        conn.execute("DELETE FROM jadwal")
        conn.execute("DELETE FROM anggota_bagian")
        conn.execute("DELETE FROM hari_libur")
        conn.execute("DELETE FROM bagian")
        conn.execute("DELETE FROM periode")
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
# =========================================================
# ALGORITMA KOMBINASI DANA
# =========================================================

def cari_beberapa_kombinasi(pegawai, dana, max_hasil=None):
    """
    Mencari beberapa kombinasi pegawai yang total tarifnya
    tepat sama dengan dana.

    Kombinasi dibedakan berdasarkan jumlah pegawai per golongan/tarif,
    bukan berdasarkan nama pegawai.
    """

    # Kelompokkan pegawai berdasarkan golongan utama + tarif
    kelompok = {}

    for i, p in enumerate(pegawai):
        golongan = p[5] or "Tidak diketahui"
        tarif = p[7]

        key = (golongan, tarif)

        if key not in kelompok:
            kelompok[key] = []

        kelompok[key].append(i)

    # Urutkan golongan supaya hasil konsisten
    kelompok_list = sorted(
        kelompok.items(),
        key=lambda x: (x[0][0], x[0][1])
    )

    hasil_komposisi = []

    def cari(index, sisa, pilihan):
        # Batasi hasil hanya jika pemanggil memang memberi batas.
        if max_hasil is not None and len(hasil_komposisi) >= max_hasil:
            return

        # Dana sudah tepat
        if sisa == 0:
            # Ubah pilihan jumlah menjadi signature
            signature = tuple(
                sorted(
                    (gol, tarif, jumlah)
                    for (gol, tarif), jumlah in pilihan.items()
                    if jumlah > 0
                )
            )

            if signature not in [
                h["signature"] for h in hasil_komposisi
            ]:
                hasil_komposisi.append({
                    "signature": signature,
                    "pilihan": pilihan.copy()
                })

            return

        # Semua kelompok sudah diperiksa
        if index >= len(kelompok_list):
            return

        (golongan, tarif), daftar_index = kelompok_list[index]

        maksimal = min(
            len(daftar_index),
            sisa // tarif
        )

        # Coba jumlah sedikit sampai banyak
        for jumlah in range(maksimal + 1):
            if jumlah > 0:
                pilihan[(golongan, tarif)] = jumlah

            cari(
                index + 1,
                sisa - (jumlah * tarif),
                pilihan
            )

            if jumlah > 0:
                del pilihan[(golongan, tarif)]

            if max_hasil is not None and len(hasil_komposisi) >= max_hasil:
                return

    cari(0, dana, {})

    # Prioritas:
    # 1. jumlah pegawai lebih sedikit
    # 2. jumlah golongan yang digunakan lebih sedikit
    hasil_komposisi.sort(
        key=lambda h: (
            sum(x[2] for x in h["signature"]),
            len(h["signature"])
        )
    )

    # Ubah komposisi menjadi daftar ID pegawai
    hasil = []

    for nomor, komposisi in enumerate(hasil_komposisi, start=1):
        pegawai_terpilih = []

        for (golongan, tarif), jumlah in komposisi["pilihan"].items():
            key = (golongan, tarif)
            daftar_index = kelompok[key]

            pegawai_terpilih.extend(
                daftar_index[:jumlah]
            )

        hasil.append({
            "nomor": nomor,
            "pegawai": pegawai_terpilih,
            "signature": komposisi["signature"]
        })

    return hasil


def cari_saran_dana_terdekat(pegawai, target):
    """Mencari jumlah tarif yang paling dekat dengan target, di bawah dan di atas."""
    target = int(target or 0)
    if target < 0:
        return None, None
    if target == 0:
        return 0, 0

    nilai = [int(p[7] or 0) for p in pegawai if int(p[7] or 0) > 0]
    if not nilai:
        return 0, None

    # Gunakan satuan terbesar yang membagi semua nilai agar DP tetap ringan.
    skala = 0
    for v in nilai + [target]:
        skala = math.gcd(skala, v)
    skala = max(skala, 1)

    vals = [v // skala for v in nilai]
    target_u = target // skala

    # Cari sedikit ruang di atas target. Karena setiap nominal positif,
    # cukup diperluas dengan nominal terbesar untuk menemukan kandidat terdekat.
    cap = target_u + max(vals)
    reachable = [False] * (cap + 1)
    reachable[0] = True

    for v in vals:
        for total in range(cap, v - 1, -1):
            if reachable[total - v]:
                reachable[total] = True

    lower_u = None
    for total in range(min(target_u, cap), -1, -1):
        if reachable[total]:
            lower_u = total
            break

    higher_u = None
    for total in range(target_u, cap + 1):
        if reachable[total]:
            higher_u = total
            break

    lower = lower_u * skala if lower_u is not None else None
    higher = higher_u * skala if higher_u is not None else None
    return lower, higher


# =========================================================
# APLIKASI
# =========================================================

class App(tk.Tk):

    def __init__(self):

        super().__init__()

        self.title(
            "Sistem Penugasan Pegawai"
        )

        self.geometry(
            "1400x900"
        )

        self.minsize(
            1200,
            700
        )
        self.state("zoomed")

        create_database()
        ensure_bagian_schema()

        self.pegawai = load_pegawai()

        now = datetime.now()

        self.bulan = now.month
        self.tahun = now.year

        self.periode_id = get_or_create_periode(
            self.bulan,
            self.tahun
        )

        kode_awal = load_bagian_codes(self.periode_id)
        self.bagian_aktif = kode_awal[0] if kode_awal else "A"

        self.selected_ids = set()

        self.bagian_id = None

        self.build_style()
        self.build_header()
        self.build_main()

        self.load_bagian_data()
        self.load_kalender()

    # =====================================================
    # STYLE
    # =====================================================

    def build_style(self):

        style = ttk.Style()

        try:
            style.theme_use("clam")
        except:
            pass

        style.configure(
            "Treeview",
            rowheight=40,
            font=("Segoe UI", 11), 
            borderwidth=1,
            relief="solid"
        )

        style.configure(
            "Treeview.Heading",
            font=("Segoe UI", 11, "bold"),
            borderwidth=1,
            relief="solid"
        )

        style.configure(
            "TButton",
            font=("Segoe UI", 9)
        )

    # =====================================================
    # HEADER
    # =====================================================

    def build_header(self):

        header = tk.Frame(
            self,
            bg="#17324D",
            height=85
        )

        header.pack(
            fill="x"
        )

        tk.Label(
            header,
            text="SISTEM PENUGASAN PEGAWAI",
            bg="#17324D",
            fg="white",
            font=("Segoe UI", 21, "bold")
        ).pack(
            pady=(13, 0)
        )

        self.header_period = tk.Label(
            header,
            text="",
            bg="#17324D",
            fg="#D9E6F2",
            font=("Segoe UI", 10)
        )

        self.header_period.pack()

    # =====================================================
    # MAIN
    # =====================================================

    def build_main(self):

        # PERIODE
        periode_frame = tk.Frame(
            self,
            bg="#F4F6F8"
        )

        periode_frame.pack(
            fill="x",
            padx=15,
            pady=10
        )

        tk.Label(
            periode_frame,
            text="Periode:",
            bg="#F4F6F8",
            font=("Segoe UI", 10, "bold")
        ).pack(
            side="left"
        )

        self.bulan_var = tk.StringVar(
            value=NAMA_BULAN[self.bulan]
        )

        self.combo_bulan = ttk.Combobox(
            periode_frame,
            textvariable=self.bulan_var,
            values=NAMA_BULAN[1:],
            state="readonly",
            width=15
        )

        self.combo_bulan.pack(
            side="left",
            padx=5
        )

        self.combo_bulan.bind(
            "<<ComboboxSelected>>",
            self.change_period
        )

        self.tahun_var = tk.StringVar(
            value=str(self.tahun)
        )

        ttk.Entry(
            periode_frame,
            textvariable=self.tahun_var,
            width=8
        ).pack(
            side="left",
            padx=5
        )

        ttk.Button(
            periode_frame,
            text="Gunakan Periode",
            command=self.change_period
        ).pack(
            side="left",
            padx=5
        )
        
        # DATA KESELURUHAN
        # Satu menu untuk membuka data master pegawai atau kalender keseluruhan.
        self.data_keseluruhan_menu = tk.Menu(
            self,
            tearoff=False
        )
        self.data_keseluruhan_menu.add_command(
            label="👥 Data Pegawai",
            command=self.buka_data_pegawai
        )
        self.data_keseluruhan_menu.add_command(
            label="📅 Kalender Keseluruhan",
            command=self.buka_kalender_keseluruhan
        )
        self.data_keseluruhan_menu.add_separator()
        self.data_keseluruhan_menu.add_command(
            label="🧹 Kelola / Bersihkan Data",
            command=self.buka_kelola_data
        )

        ttk.Menubutton(
            periode_frame,
            text="Data Keseluruhan ▾",
            menu=self.data_keseluruhan_menu
        ).pack(side="left", padx=4)

        ttk.Button(
            periode_frame,
            text="📤 Export Excel",
            command=self.export_penugasan_excel
        ).pack(side="left", padx=4)

        ttk.Button(
            periode_frame,
            text="↻ Refresh Semua",
            command=self.refresh_all
        ).pack(side="left", padx=4)

        # DAFTAR KEGIATAN DINAMIS
        self.bagian_frame = tk.Frame(
            self,
            bg="#F4F6F8"
        )
        self.bagian_frame.pack(
            fill="x",
            padx=15,
            pady=5
        )

        tk.Label(
            self.bagian_frame,
            text="Kegiatan:",
            bg="#F4F6F8",
            font=("Segoe UI", 10, "bold")
        ).pack(side="left")

        self.bagian_buttons = {}
        self.refresh_bagian_buttons()

        # NOTEBOOK
        self.notebook = ttk.Notebook(
            self
        )

        self.notebook.pack(
            fill="both",
            expand=True,
            padx=15,
            pady=10
        )

        self.tab_pengaturan = tk.Frame(
            self.notebook,
            bg="#F4F6F8"
        )

        self.tab_jadwal = tk.Frame(
            self.notebook,
            bg="#F4F6F8"
        )

        self.tab_data_kegiatan = tk.Frame(
            self.notebook,
            bg="#F4F6F8"
        )

        self.tab_pegawai = tk.Frame(
            self.notebook,
            bg="#F4F6F8"
        )

        self.notebook.add(
            self.tab_pengaturan,
            text="  Pengaturan Kegiatan  "
        )

        self.notebook.add(
            self.tab_jadwal,
            text="  Kalender Kegiatan  "
        )

        self.notebook.add(
            self.tab_data_kegiatan,
            text="  Data Kegiatan  "
        )

        self.build_pengaturan()
        self.build_jadwal()
        self.build_data_kegiatan()

        # Data Kegiatan selalu disegarkan saat tab dibuka/aktif kembali.
        # Jadi user tidak perlu menekan Refresh Semua hanya untuk melihat
        # perubahan terbaru pada kegiatan yang sedang dipilih.
        self.notebook.bind(
            "<<NotebookTabChanged>>",
            self.on_notebook_tab_changed
        )

    def on_notebook_tab_changed(self, event=None):
        """Refresh tampilan tab aktif yang datanya berasal dari database."""
        try:
            current = self.notebook.select()
            if current == str(self.tab_data_kegiatan):
                self.refresh_data_kegiatan()
            elif current == str(self.tab_jadwal):
                self.load_kalender()
            elif current == str(self.tab_pengaturan):
                self.refresh_selection_table()
        except (tk.TclError, AttributeError):
            pass

    def refresh_bagian_buttons(self):
        """Membangun tombol kegiatan dan tombol X hapus untuk tiap kegiatan."""
        for widget in self.bagian_frame.winfo_children()[1:]:
            widget.destroy()
        self.bagian_buttons = {}

        conn = get_connection()
        rows = conn.execute(
            "SELECT kode, COALESCE(nama_bagian, '') FROM bagian WHERE periode_id = ? ORDER BY id",
            (self.periode_id,)
        ).fetchall()
        conn.close()

        for kode, nama_bagian in rows:
            wrap = tk.Frame(self.bagian_frame, bg="#F4F6F8")
            wrap.pack(side="left", padx=3)

            judul = (nama_bagian or "").strip() or f"Kegiatan {kode}"
            btn = tk.Button(
                wrap,
                text=judul,
                command=lambda k=kode: self.change_bagian(k),
                width=16,
                font=("Segoe UI", 9, "bold")
            )
            btn.pack(side="left")
            btn.bind(
                "<Double-Button-1>",
                lambda event, k=kode: self.rename_bagian(k)
            )

            btn_x = tk.Button(
                wrap,
                text="×",
                command=lambda k=kode: self.hapus_data_bagian(k),
                width=2,
                font=("Segoe UI", 9, "bold"),
                padx=0,
                pady=0
            )
            btn_x.pack(side="left", padx=(1, 0))

            self.bagian_buttons[kode] = btn

        ttk.Button(
            self.bagian_frame,
            text="＋ Tambah Kegiatan",
            command=self.tambah_bagian
        ).pack(side="left", padx=(10, 4))

        self.update_bagian_button_state()

    def update_bagian_button_state(self):
        for kode, btn in self.bagian_buttons.items():
            if kode == self.bagian_aktif:
                btn.config(bg="#17324D", fg="white")
            else:
                btn.config(bg="#E0E6EC", fg="black")

    def rename_bagian(self, kode):
        """Rename kegiatan dengan double-click pada nama tombol kegiatan."""
        data = load_bagian(self.periode_id, kode)
        if not data:
            return

        bagian_id = data[0]
        conn = get_connection()
        row = conn.execute(
            "SELECT COALESCE(nama_bagian, '') FROM bagian WHERE id = ?",
            (bagian_id,)
        ).fetchone()
        conn.close()

        nama_lama = (row[0] if row else "").strip() or f"Kegiatan {kode}"
        nama_baru = simpledialog.askstring(
            "Ganti Nama Kegiatan",
            f"Nama baru untuk {nama_lama}:",
            initialvalue=nama_lama,
            parent=self
        )

        if nama_baru is None:
            return

        nama_baru = nama_baru.strip()
        if not nama_baru:
            messagebox.showwarning(
                "Nama Kegiatan",
                "Nama kegiatan tidak boleh kosong."
            )
            return

        conn = None
        try:
            conn = get_connection()
            conn.execute(
                "UPDATE bagian SET nama_bagian = ? WHERE id = ?",
                (nama_baru, bagian_id)
            )
            conn.commit()
        except Exception as e:
            if conn:
                conn.rollback()
            messagebox.showerror("Gagal Mengganti Nama", str(e))
            return
        finally:
            if conn:
                conn.close()

        self.refresh_bagian_buttons()
        self.load_bagian_data()

    # Double-click nama kegiatan = rename.
    def tambah_bagian(self):
        kode_list = load_bagian_codes(self.periode_id)
        kode_baru = kode_bagian_berikutnya(kode_list)

        conn = None
        try:
            conn = get_connection()
            cursor = conn.cursor()
            cursor.execute("""
                INSERT INTO bagian
                (periode_id, kode, nama_bagian, kegiatan, lokasi, dana, penanggung_jawab_id)
                VALUES (?, ?, ?, '', '', 0, NULL)
            """, (self.periode_id, kode_baru, f"Kegiatan {kode_baru}"))
            conn.commit()
        except sqlite3.IntegrityError:
            if conn:
                conn.rollback()
            # Ambil ulang kode karena data mungkin berubah saat proses berlangsung.
            kode_list = load_bagian_codes(self.periode_id)
            kode_baru = kode_bagian_berikutnya(kode_list)
            try:
                cursor.execute("""
                    INSERT INTO bagian
                    (periode_id, kode, nama_bagian, kegiatan, lokasi, dana, penanggung_jawab_id)
                    VALUES (?, ?, ?, '', '', 0, NULL)
                """, (self.periode_id, kode_baru, f"Kegiatan {kode_baru}"))
                conn.commit()
            except Exception:
                conn.rollback()
                messagebox.showerror("Gagal", "Kegiatan baru gagal ditambahkan. Silakan coba lagi.")
                return
        except sqlite3.OperationalError as e:
            if conn:
                conn.rollback()
            if "locked" in str(e).lower():
                messagebox.showerror(
                    "Database Sedang Dipakai",
                    "Database sedang terkunci. Tutup program/Excel lain yang sedang membuka database, lalu coba lagi."
                )
            else:
                messagebox.showerror("Gagal", f"Kegiatan baru gagal ditambahkan.\n\n{e}")
            return
        finally:
            if conn:
                conn.close()

        self.bagian_aktif = kode_baru
        self.refresh_bagian_buttons()
        self.load_bagian_data()
        self.load_kalender()

        self.notebook.select(self.tab_pengaturan)
        messagebox.showinfo(
            "Kegiatan Ditambahkan",
            f"Kegiatan {kode_baru} berhasil ditambahkan. Silakan isi nama dan data kegiatannya."
        )

    # =====================================================
    # Refresh ALL
    # =====================================================
    def refresh_all(self):

        try:

            # Reload pegawai dari database
            self.pegawai = load_pegawai()

            # Pastikan periode tetap sesuai pilihan saat ini
            bulan = NAMA_BULAN.index(
                self.bulan_var.get()
            )

            tahun = int(
                self.tahun_var.get()
            )

            self.bulan = bulan
            self.tahun = tahun

            self.periode_id = get_or_create_periode(
                bulan,
                tahun
            )

            kode_list = load_bagian_codes(self.periode_id)
            if self.bagian_aktif not in kode_list:
                self.bagian_aktif = kode_list[0]
            self.refresh_bagian_buttons()

            # Reload kegiatan aktif
            self.load_bagian_data()

            # Reload kalender
            self.load_kalender()

            # Reload tabel data pegawai
            if hasattr(
                self,
                "refresh_data_pegawai"
            ):
                self.refresh_data_pegawai()

            # Refresh seluruh tampilan yang bergantung pada database.
            # Satu pintu refresh supaya setelah import/edit/hapus tidak ada
            # tampilan yang tertinggal memakai data lama.
            self.refresh_pj_options()
            self.refresh_selection_table()
            self.refresh_data_kegiatan()
            self.load_kalender_keseluruhan()

            # Update header
            nama_header = self.nama_bagian_var.get().strip() if hasattr(self, "nama_bagian_var") else f"Kegiatan {self.bagian_aktif}"
            self.header_period.config(
                text=(
                    f"{NAMA_BULAN[self.bulan]} "
                    f"{self.tahun}"
                    f"  •  {nama_header}"
                )
            )

        except Exception as e:

            messagebox.showerror(
                "Refresh Gagal",
                f"Terjadi kesalahan saat refresh:\n\n{e}"
            )

    # =====================================================
    # PENGATURAN
    # =====================================================

    def build_pengaturan(self):

        # Nama kegiatan hanya disimpan sebagai state internal.
        # Tidak ditampilkan sebagai field terpisah; rename dilakukan
        # dengan double-click pada tombol kegiatan di bagian atas.
        self.nama_bagian_var = tk.StringVar(value="")

        # =====================================================
        # CANVAS SCROLLABLE
        # =====================================================

        container = tk.Frame(
            self.tab_pengaturan,
            bg="#F4F6F8"
        )
        container.pack(
            fill="both",
            expand=True
        )

        canvas = tk.Canvas(
            container,
            bg="#F4F6F8",
            highlightthickness=0
        )
        canvas.pack(
            side="left",
            fill="both",
            expand=True
        )

        scrollbar = ttk.Scrollbar(
            container,
            orient="vertical",
            command=canvas.yview
        )
        scrollbar.pack(
            side="right",
            fill="y"
        )

        canvas.configure(
            yscrollcommand=scrollbar.set
        )

        # Frame isi yang bisa discroll
        scroll_frame = tk.Frame(
            canvas,
            bg="#F4F6F8"
        )

        canvas_window = canvas.create_window(
            (0, 0),
            window=scroll_frame,
            anchor="nw"
        )

        def update_scroll_region(event=None):
            canvas.configure(
                scrollregion=canvas.bbox("all")
            )

        def resize_scroll_frame(event):
            canvas.itemconfig(
                canvas_window,
                width=event.width
            )

        scroll_frame.bind(
            "<Configure>",
            update_scroll_region
        )

        canvas.bind(
            "<Configure>",
            resize_scroll_frame
        )

        # Mouse wheel
        def scroll_mousewheel(event):
            canvas.yview_scroll(
                int(-1 * (event.delta / 120)),
                "units"
            )

        canvas.bind_all(
            "<MouseWheel>",
            scroll_mousewheel
        )

        # =====================================================
        # LAYOUT UTAMA: PENGATURAN DI KIRI, PEGAWAI DI KANAN
        # =====================================================
        frame = tk.Frame(scroll_frame, bg="#F4F6F8")
        frame.pack(fill="both", expand=True, padx=20, pady=18)
        frame.grid_columnconfigure(0, weight=0, minsize=390)
        frame.grid_columnconfigure(1, weight=1)
        # Baris atas hanya untuk pengaturan dan kontrol pegawai.
        # Tabel pegawai ditempatkan di baris kedua agar dapat memakai
        # seluruh lebar layar, termasuk ruang kosong di bawah pengaturan.
        frame.grid_rowconfigure(0, weight=0)
        frame.grid_rowconfigure(1, weight=1)

        settings = tk.LabelFrame(
            frame, text="  PENGATURAN KEGIATAN  ", bg="#F4F6F8",
            font=("Segoe UI", 12, "bold"), padx=18, pady=14
        )
        settings.grid(row=0, column=0, sticky="nsew", padx=(0, 14))
        settings.grid_columnconfigure(1, weight=1)

        tk.Label(settings, text="Deskripsi", bg="#F4F6F8").grid(row=0, column=0, sticky="nw", pady=8)
        # State teks deskripsi. StringVar dipertahankan untuk kompatibilitas
        # dengan loader lama, sedangkan widget utama menggunakan Text agar
        # deskripsi panjang bisa dibaca dan diedit beberapa baris.
        self.kegiatan_var = tk.StringVar(value="")
        self.kegiatan_text = tk.Text(
            settings, height=4, width=34, wrap="word",
            font=("Segoe UI", 10), relief="solid", borderwidth=1
        )
        self.kegiatan_text.grid(row=0, column=1, sticky="ew", pady=8)

        tk.Label(settings, text="Lokasi", bg="#F4F6F8").grid(row=1, column=0, sticky="w", pady=8)
        self.lokasi_var = tk.StringVar()
        ttk.Entry(settings, textvariable=self.lokasi_var, width=34).grid(row=1, column=1, sticky="ew", pady=8)

        tk.Label(settings, text="Penanggung Jawab", bg="#F4F6F8", font=("Segoe UI", 10, "bold")).grid(row=2, column=0, sticky="w", pady=8)
        self.pj_var = tk.StringVar()
        self.pj_options = {}
        self.combo_pj = ttk.Combobox(settings, textvariable=self.pj_var, state="readonly", width=31)
        self.combo_pj.grid(row=2, column=1, sticky="ew", pady=8)

        tk.Label(settings, text="Dana Bulanan", bg="#F4F6F8", font=("Segoe UI", 10, "bold")).grid(row=3, column=0, sticky="w", pady=8)
        self.dana_var = tk.StringVar()
        ttk.Entry(settings, textvariable=self.dana_var, width=20).grid(row=3, column=1, sticky="w", pady=8)

        tk.Label(settings, text="Double-click nama kegiatan di atas untuk mengganti nama.", bg="#F4F6F8", fg="#666666", font=("Segoe UI", 9)).grid(row=4, column=0, columnspan=2, sticky="w", pady=(14, 4))
        tk.Label(settings, text="Nama kegiatan hanya label tampilan.", bg="#F4F6F8", fg="#777777", font=("Segoe UI", 8)).grid(row=5, column=0, columnspan=2, sticky="w")

        # =====================================================
        # KONTROL PEGAWAI (BARIS ATAS)
        # =====================================================
        # Panel ini hanya berisi instruksi, tombol otomatis, total, simpan,
        # dan pencarian. Tabel pegawai dipindahkan ke baris kedua agar
        # memanfaatkan seluruh lebar jendela.
        pegawai_panel = tk.LabelFrame(
            frame, text="  PEGAWAI UNTUK KEGIATAN INI  ", bg="#F4F6F8",
            font=("Segoe UI", 12, "bold"), padx=10, pady=8
        )
        pegawai_panel.grid(row=0, column=1, sticky="nsew")
        pegawai_panel.grid_columnconfigure(0, weight=1)

        tk.Label(
            pegawai_panel,
            text="Pilih pegawai secara manual, atau gunakan penentuan otomatis berdasarkan dana.",
            bg="#F4F6F8", fg="#666666"
        ).grid(row=0, column=0, sticky="w", pady=(0, 7))

        tombol_frame = tk.Frame(pegawai_panel, bg="#F4F6F8")
        tombol_frame.grid(row=1, column=0, sticky="w", pady=(0, 7))
        ttk.Button(
            tombol_frame,
            text="AUTO TENTUKAN BERDASARKAN DANA",
            command=self.auto_tentukan
        ).pack(side="left", padx=(0, 6))
        ttk.Button(
            tombol_frame,
            text="⚙ ATUR KOMPOSISI",
            command=self.buka_komposisi_manual
        ).pack(side="left", padx=(0, 6))
        self.btn_acak_lagi = ttk.Button(
            tombol_frame,
            text="🔀 ACAK LAGI",
            command=self.acak_lagi,
            state="disabled"
        )
        self.btn_acak_lagi.pack(side="left")

        self.info_dana = tk.Label(
            pegawai_panel,
            text="Pegawai terpilih: 0    •    Total: Rp 0",
            bg="#F4F6F8",
            font=("Segoe UI", 11, "bold")
        )
        self.info_dana.grid(row=2, column=0, sticky="w", pady=(2, 7))

        # Simpan tetap berada di bagian atas agar mudah ditemukan.
        save_frame = tk.Frame(pegawai_panel, bg="#F4F6F8")
        save_frame.grid(row=3, column=0, sticky="ew", pady=(0, 7))
        save_frame.grid_columnconfigure(0, weight=1)
        ttk.Button(
            save_frame,
            text="💾 Simpan Pengaturan",
            command=self.simpan_bagian
        ).grid(row=0, column=0, sticky="e", ipadx=18, ipady=7)

        # =====================================================
        # TABEL PEGAWAI FULL WIDTH
        # =====================================================
        self.build_tabel_pilih(frame)

    def build_tabel_pilih(self, parent):

        # Area pegawai dibagi dua: tabel di kiri dan ringkasan komposisi
        # yang dipilih di kanan. Ini membantu pengguna memilih manual.
        selection_area = tk.Frame(
            parent,
            bg="#F4F6F8",
            height=520
        )
        selection_area.grid_propagate(False)

        selection_area.grid(
            row=1,
            column=0,
            columnspan=2,
            sticky="nsew",
            pady=(12, 10)
        )
        selection_area.grid_columnconfigure(0, weight=1)
        selection_area.grid_rowconfigure(0, weight=0)
        selection_area.grid_rowconfigure(1, weight=1)

        # Pencarian ditempatkan tepat di atas daftar nama pegawai, bukan
        # di dalam panel kontrol "Pegawai untuk Kegiatan Ini".
        search_frame = tk.Frame(selection_area, bg="#F4F6F8")
        search_frame.grid(
            row=0, column=0, sticky="ew",
            padx=(0, 0), pady=(0, 8)
        )
        search_frame.grid_columnconfigure(1, weight=1)

        tk.Label(
            search_frame,
            text="Cari Pegawai:",
            bg="#F4F6F8",
            font=("Segoe UI", 10, "bold")
        ).grid(row=0, column=0, sticky="w", padx=(0, 8))

        self.search_pegawai_var = tk.StringVar()
        self.entry_search_pegawai = ttk.Entry(
            search_frame,
            textvariable=self.search_pegawai_var
        )
        self.entry_search_pegawai.grid(row=0, column=1, sticky="ew")
        self.entry_search_pegawai.bind(
            "<KeyRelease>",
            lambda event: self.refresh_selection_table()
        )
        ttk.Button(
            search_frame,
            text="✕",
            width=3,
            command=lambda: (
                self.search_pegawai_var.set(""),
                self.refresh_selection_table()
            )
        ).grid(row=0, column=2, padx=(5, 0))

        table_area = tk.Frame(
            selection_area,
            bg="#F4F6F8"
        )
        table_area.grid(row=1, column=0, sticky="nsew")
        table_area.grid_columnconfigure(0, weight=1)
        table_area.grid_columnconfigure(1, weight=0)
        table_area.grid_rowconfigure(0, weight=1)

        table_frame = tk.Frame(
            table_area,
            bg="white",
            height=520
        )

        table_frame.grid(row=0, column=0, sticky="nsew")
        table_frame.pack_propagate(False)

        columns = (
            "pilih",
            "nama",
            "golongan",
            "jabatan",
            "tarif"
        )

        self.table_pilih = ttk.Treeview(
            table_frame,
            columns=columns,
            show="headings",
            selectmode="none"
        )

        heads = {
            "pilih": "✓",
            "nama": "Nama",
            "golongan": "Golongan",
            "jabatan": "Jabatan",
            "tarif": "Tarif"
        }

        widths = {
            "pilih": 60,
            "nama": 330,
            "golongan": 110,
            "jabatan": 260,
            "tarif": 120
        }

        for col in columns:
            if col == "pilih":
                self.table_pilih.heading(
                    col,
                    text="✓",
                    command=self.toggle_semua_pegawai
                )
            else:
                self.table_pilih.heading(
                    col,
                    text=heads[col]
                )

            self.table_pilih.column(
                col,
                width=widths[col],
                anchor="center"
            )

        self.table_pilih.column("nama", anchor="w")
        self.table_pilih.column("jabatan", anchor="w")

        self.table_pilih.pack(
            side="left",
            fill="both",
            expand=True
        )

        scroll_table = ttk.Scrollbar(
            table_frame,
            orient="vertical",
            command=self.table_pilih.yview
        )
        scroll_table.pack(side="right", fill="y")

        self.table_pilih.configure(yscrollcommand=scroll_table.set)
        self.table_pilih.bind("<ButtonRelease-1>", self.toggle_pegawai)

        # Ringkasan komposisi manual di sebelah kanan.
        self.komposisi_frame = tk.Frame(
            table_area,
            bg="white",
            width=340,
            height=520,
            highlightthickness=1,
            highlightbackground="#D5DADF"
        )
        self.komposisi_frame.grid(
            row=0, column=1, sticky="nsew",
            padx=(12, 0)
        )
        table_area.grid_columnconfigure(1, minsize=340)
        self.komposisi_frame.pack_propagate(False)

        tk.Label(
            self.komposisi_frame,
            text="KOMPOSISI TERPILIH",
            bg="white",
            font=("Segoe UI", 11, "bold")
        ).pack(anchor="w", padx=15, pady=(14, 3))

        tk.Label(
            self.komposisi_frame,
            text="Ringkasan pegawai terpilih",
            bg="white",
            fg="#666666",
            font=("Segoe UI", 9)
        ).pack(anchor="w", padx=15)

        self.komposisi_text = tk.Text(
            self.komposisi_frame,
            height=18,
            width=32,
            bg="white",
            relief="flat",
            borderwidth=0,
            font=("Segoe UI", 10),
            state="disabled"
        )
        self.komposisi_text.pack(
            fill="both",
            expand=True,
            padx=12,
            pady=(8, 5)
        )

        self.refresh_selection_table()

    def refresh_komposisi_manual(self):
        """Bangun ulang panel komposisi dari nol berdasarkan pilihan TERKINI."""
        if not hasattr(self, "komposisi_text"):
            return

        # Selalu bersihkan isi lama terlebih dahulu. Ini penting saat
        # komposisi/kegiatan diganti agar teks dari pilihan sebelumnya
        # tidak pernah tertinggal di panel.
        self.komposisi_text.configure(state="normal")
        self.komposisi_text.delete("1.0", "end")

        kelompok = {}
        total = 0
        jumlah = 0

        selected_ids = set(getattr(self, "selected_ids", set()))
        for p in self.pegawai:
            if p[0] not in selected_ids:
                continue

            # Gunakan GOLONGAN UTAMA (p[5]), bukan golongan asli p[4].
            golongan = (p[5] or "-").strip()
            tarif = int(p[7] or 0)
            key = (golongan, tarif)
            kelompok[key] = kelompok.get(key, 0) + 1
            total += tarif
            jumlah += 1

        lines = []
        if not kelompok:
            lines.append("Belum ada pegawai yang dipilih.")
        else:
            # Penanggung jawab selalu wajib ikut terpilih. Tampilkan terpisah
            # supaya golongan PJ tidak terlihat seperti sisa dari komposisi lama.
            pj_id = getattr(self, "penanggung_jawab_id", None)
            pj = next((p for p in self.pegawai if p[0] == pj_id and p[0] in selected_ids), None)
            if pj:
                lines.append(f"PJ wajib: {pj[1]}")
                lines.append(f"• Gol {((pj[5] or '-').strip())}: {rupiah(pj[7] or 0)}")
                lines.append("")

            # Tampilkan hanya berdasarkan pilihan TERKINI.
            for (golongan, tarif), banyak in sorted(
                kelompok.items(),
                key=lambda x: (str(x[0][0]), x[0][1])
            ):
                subtotal = banyak * tarif
                lines.append(f"• Gol {golongan}: {banyak} orang")
                lines.append(
                    f"  {rupiah(tarif)} × {banyak} = {rupiah(subtotal)}"
                )

            lines.append("")
            lines.append(f"Jumlah pegawai : {jumlah} orang")
            lines.append(f"Total tarif    : {rupiah(total)}")

        self.komposisi_text.insert("1.0", "\n".join(lines))
        self.komposisi_text.configure(state="disabled")

    def pilih_kombinasi_dialog(self, hasil, dana, nama_pj):
        """Dialog pilihan kombinasi dengan kartu yang bisa membungkus teks."""
        win = tk.Toplevel(self)
        win.title("Pilih Komposisi Kegiatan")
        win.geometry("1200x760")
        win.minsize(950, 620)
        win.transient(self)
        win.grab_set()

        tk.Label(
            win,
            text=f"Ditemukan {len(hasil)} kombinasi yang tepat untuk {rupiah(dana)}",
            font=("Segoe UI", 14, "bold")
        ).pack(anchor="w", padx=24, pady=(18, 4))
        tk.Label(
            win,
            text=f"PJ wajib: {nama_pj}  •  Pilih salah satu komposisi di bawah.",
            fg="#666666", font=("Segoe UI", 10)
        ).pack(anchor="w", padx=24, pady=(0, 12))

        outer = tk.Frame(win, bg="#E9EDF0")
        outer.pack(fill="both", expand=True, padx=24, pady=(0, 12))
        outer.grid_rowconfigure(0, weight=1)
        outer.grid_columnconfigure(0, weight=1)

        canvas = tk.Canvas(outer, highlightthickness=0, bg="#E9EDF0")
        canvas.grid(row=0, column=0, sticky="nsew")
        scrollbar = ttk.Scrollbar(outer, orient="vertical", command=canvas.yview)
        scrollbar.grid(row=0, column=1, sticky="ns")
        canvas.configure(yscrollcommand=scrollbar.set)

        content = tk.Frame(canvas, bg="#E9EDF0")
        window_id = canvas.create_window((0, 0), window=content, anchor="nw")
        content.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.bind("<Configure>", lambda e: canvas.itemconfigure(window_id, width=e.width))

        selected = tk.IntVar(value=hasil[0]["nomor"] if hasil else 0)

        def format_signature(signature):
            parts = []
            for golongan, tarif, jumlah in signature:
                parts.append(f"Gol {golongan}: {jumlah} orang ({rupiah(tarif)} / orang)")
            return "  •  ".join(parts)

        for item in hasil:
            card = tk.Frame(content, bg="white", bd=1, relief="solid", padx=12, pady=10)
            card.pack(fill="x", padx=8, pady=5)
            signature_text = format_signature(item["signature"])
            jumlah_total = sum(x[2] for x in item["signature"])
            total = sum(x[1] * x[2] for x in item["signature"])

            rb = ttk.Radiobutton(card, variable=selected, value=item["nomor"])
            rb.grid(row=0, column=0, rowspan=2, padx=(2, 10), sticky="n")
            tk.Label(card, text=f"Komposisi {item['nomor']}", bg="white", font=("Segoe UI", 10, "bold")).grid(row=0, column=1, sticky="w")
            tk.Label(card, text=signature_text, bg="white", justify="left", anchor="w", wraplength=820, font=("Segoe UI", 10)).grid(row=1, column=1, sticky="w", pady=(4, 0))
            tk.Label(card, text=f"{jumlah_total} orang  •  {rupiah(total)}", bg="white", fg="#555555", font=("Segoe UI", 9, "bold")).grid(row=0, column=2, rowspan=2, padx=(18, 4), sticky="e")
            card.grid_columnconfigure(1, weight=1)
            card.bind("<Button-1>", lambda e, n=item["nomor"]: selected.set(n))

        result = {"pilihan": None}
        buttons = tk.Frame(win)
        buttons.pack(fill="x", padx=24, pady=(0, 18))

        def pilih():
            if not selected.get():
                messagebox.showwarning("Pilih Komposisi", "Pilih salah satu komposisi terlebih dahulu.", parent=win)
                return
            result["pilihan"] = selected.get()
            win.destroy()

        ttk.Button(buttons, text="Batal", command=win.destroy).pack(side="right", padx=(8, 0), ipadx=14, ipady=4)
        ttk.Button(buttons, text="Pilih Komposisi", command=pilih).pack(side="right", ipadx=14, ipady=4)
        win.wait_window()
        return result["pilihan"]

    def buka_komposisi_manual(self):
        """Atur komposisi berdasarkan golongan atau jabatan, lalu acak pegawai sesuai aturan."""
        pj_id = self.get_selected_pj_id()
        if not pj_id:
            messagebox.showwarning("Penanggung Jawab", "Pilih penanggung jawab terlebih dahulu.")
            return
        self.penanggung_jawab_id = pj_id

        gol_values = sorted({str(p[5] or "").strip() for p in self.pegawai if str(p[5] or "").strip()})
        jab_values = sorted({str(p[6] or "").strip() for p in self.pegawai if str(p[6] or "").strip()})
        criteria_values = {"Golongan": gol_values, "Jabatan": jab_values}

        win = tk.Toplevel(self)
        win.title("Atur Komposisi Manual")
        win.geometry("760x560")
        win.minsize(680, 480)
        win.transient(self)
        win.grab_set()

        tk.Label(win, text="Atur Komposisi Pegawai", font=("Segoe UI", 14, "bold")).pack(anchor="w", padx=20, pady=(18, 3))
        tk.Label(win, text="Tentukan kriteria dan jumlah pegawai. Penanggung Jawab selalu ikut terpilih.", fg="#666666").pack(anchor="w", padx=20, pady=(0, 12))

        rows_frame = tk.Frame(win)
        rows_frame.pack(fill="both", expand=True, padx=20)
        rows_frame.grid_columnconfigure(1, weight=1)

        rows = []

        def add_row(default_type="Golongan"):
            r = len(rows)
            tipe = tk.StringVar(value=default_type)
            nilai = tk.StringVar()
            jumlah = tk.IntVar(value=1)
            cb_tipe = ttk.Combobox(rows_frame, textvariable=tipe, values=("Golongan", "Jabatan"), state="readonly", width=12)
            cb_nilai = ttk.Combobox(rows_frame, textvariable=nilai, state="readonly", width=48)
            sp_jumlah = tk.Spinbox(rows_frame, from_=1, to=99, textvariable=jumlah, width=6)
            cb_tipe.grid(row=r, column=0, padx=(0, 8), pady=6, sticky="w")
            cb_nilai.grid(row=r, column=1, padx=(0, 8), pady=6, sticky="ew")
            sp_jumlah.grid(row=r, column=2, pady=6, sticky="w")
            rows.append((tipe, nilai, jumlah, cb_nilai))

            def update_values(event=None):
                vals = criteria_values.get(tipe.get(), [])
                cb_nilai.configure(values=vals)
                if nilai.get() not in vals:
                    nilai.set(vals[0] if vals else "")
            cb_tipe.bind("<<ComboboxSelected>>", update_values)
            update_values()

        tk.Label(rows_frame, text="Kriteria", font=("Segoe UI", 9, "bold")).grid(row=0, column=0, sticky="w")
        tk.Label(rows_frame, text="Pilihan", font=("Segoe UI", 9, "bold")).grid(row=0, column=1, sticky="w")
        tk.Label(rows_frame, text="Jumlah", font=("Segoe UI", 9, "bold")).grid(row=0, column=2, sticky="w")
        add_row()

        ttk.Button(win, text="＋ Tambah Kriteria", command=add_row).pack(anchor="w", padx=20, pady=10)

        preview = tk.Label(win, text="", justify="left", anchor="w", fg="#555555")
        preview.pack(fill="x", padx=20, pady=(0, 8))

        def apply_manual():
            rules = []
            for tipe, nilai, jumlah, _ in rows:
                v = nilai.get().strip()
                try:
                    n = int(jumlah.get())
                except Exception:
                    n = 0
                if not v or n <= 0:
                    continue
                rules.append((tipe.get(), v, n))
            if not rules:
                messagebox.showwarning("Komposisi Kosong", "Masukkan minimal satu kriteria dan jumlah pegawai.", parent=win)
                return

            selected = self.pilih_pegawai_manual_berdasarkan_aturan(rules, pj_id)
            if selected is None:
                return
            self.selected_ids = {self.pegawai[i][0] for i in selected}
            self.komposisi_aktif = None
            self.refresh_selection_table()
            self.refresh_komposisi_manual()
            if hasattr(self, "btn_acak_lagi"):
                self.btn_acak_lagi.config(state="disabled")
            win.destroy()

        bottom = tk.Frame(win)
        bottom.pack(fill="x", padx=20, pady=(0, 18))
        ttk.Button(bottom, text="Batal", command=win.destroy).pack(side="right", padx=(8, 0))
        ttk.Button(bottom, text="🎲 Acak Sesuai Komposisi", command=apply_manual).pack(side="right")

    def pilih_pegawai_manual_berdasarkan_aturan(self, rules, pj_id):
        """Memilih pegawai acak sesuai aturan Golongan/Jabatan, dengan PJ wajib."""
        riwayat = load_riwayat_turun(self.periode_id)
        pool = [i for i, p in enumerate(self.pegawai) if p[0] != pj_id]
        selected = []
        used = set()

        for tipe, nilai, jumlah in rules:
            if tipe == "Golongan":
                candidates = [i for i in pool if i not in used and str(self.pegawai[i][5] or "").strip() == nilai]
            else:
                candidates = [i for i in pool if i not in used and str(self.pegawai[i][6] or "").strip() == nilai]
            if len(candidates) < jumlah:
                messagebox.showwarning(
                    "Pegawai Tidak Cukup",
                    f"Kriteria {tipe} '{nilai}' membutuhkan {jumlah} orang, tetapi hanya tersedia {len(candidates)} orang yang belum terpilih."
                )
                return None
            random.shuffle(candidates)
            candidates.sort(key=lambda i: (riwayat.get(self.pegawai[i][0], {}).get("bulan_ini", 0), riwayat.get(self.pegawai[i][0], {}).get("total", 0)))
            picked = candidates[:jumlah]
            selected.extend(picked)
            used.update(picked)

        pj_index = next((i for i, p in enumerate(self.pegawai) if p[0] == pj_id), None)
        if pj_index is None:
            messagebox.showerror("Penanggung Jawab", "Data penanggung jawab tidak ditemukan.")
            return None
        if pj_index not in selected:
            selected.insert(0, pj_index)
        return selected

    def refresh_selection_table(self):

        # Tabel belum dibuat, jangan refresh dulu
        if not hasattr(self, "table_pilih"):
            return

        # Bersihkan tabel
        for item in self.table_pilih.get_children():
            self.table_pilih.delete(item)

        # Warna baris selang-seling
        self.table_pilih.tag_configure(
            "baris1",
            background="#FFFFFF"
        )

        self.table_pilih.tag_configure(
            "baris2",
            background="#F1F4F7"
        )

        # Ambil teks pencarian
        keyword = ""

        if hasattr(
            self,
            "search_pegawai_var"
        ):
            keyword = (
                self.search_pegawai_var
                .get()
                .strip()
                .lower()
            )

        total = 0
        nomor_baris = 0

        for p in self.pegawai:

            pegawai_id = p[0]

            # =================================================
            # FILTER SEARCH
            # =================================================

            teks_pencarian = " ".join([
                str(p[1] or ""),  # Nama
                str(p[2] or ""),  # NIP
                str(p[3] or ""),  # Pangkat
                str(p[4] or ""),  # Golongan asli
                str(p[5] or ""),  # Golongan utama
                str(p[6] or "")   # Jabatan
            ]).lower()

            if keyword and keyword not in teks_pencarian:
                continue

            # =================================================
            # CHECKBOX
            # =================================================

            tanda = (
                "✓"
                if pegawai_id in self.selected_ids
                else ""
            )

            tag = (
                "baris1"
                if nomor_baris % 2 == 0
                else "baris2"
            )

            self.table_pilih.insert(
                "",
                "end",
                iid=str(pegawai_id),
                values=(
                    tanda,
                    p[1],
                    p[5],
                    p[6],
                    rupiah(p[7])
                ),
                tags=(tag,)
            )

            nomor_baris += 1

            # =================================================
            # TOTAL
            # =================================================

            if pegawai_id in self.selected_ids:
                total += p[7]

        # =====================================================
        # INFO TOTAL
        # =====================================================

        self.info_dana.config(
            text=(
                f"Pegawai terpilih: "
                f"{len(self.selected_ids)}"
                f"    •    Total: "
                f"{rupiah(total)}"
            )
        )

        self.refresh_komposisi_manual()

    # =====================================================
    # PILIH PEGAWAI MANUAL
    # =====================================================
    def toggle_semua_pegawai(self):

        semua_id = {
            p[0]
            for p in self.pegawai
        }

        # Kalau semuanya sudah dipilih → hapus semua, tetapi PJ tetap wajib.
        if self.selected_ids == semua_id:

            self.selected_ids.clear()
            if getattr(self, "penanggung_jawab_id", None):
                self.selected_ids.add(self.penanggung_jawab_id)

        # Kalau belum semuanya → pilih semua
        else:

            self.selected_ids = semua_id.copy()

        self.refresh_selection_table()

    def toggle_pegawai(self, event):

        item = self.table_pilih.identify_row(
            event.y
        )

        if not item:
            return

        pegawai_id = int(item)

        if pegawai_id == getattr(self, "penanggung_jawab_id", None):
            messagebox.showwarning(
                "Penanggung Jawab Wajib",
                "Penanggung jawab kegiatan wajib tetap dipilih."
            )
            return

        if pegawai_id in self.selected_ids:

            self.selected_ids.remove(
                pegawai_id
            )

        else:

            self.selected_ids.add(
                pegawai_id
            )

        self.refresh_selection_table()

    # =====================================================
    # SIMPAN BAGIAN
    # =====================================================

    def simpan_bagian(self):

        try:
            dana = parse_rupiah(self.dana_var.get())
        except Exception:
            messagebox.showerror("Dana", "Dana harus berupa angka.")
            return

        pj_id = self.get_selected_pj_id()
        if not pj_id:
            messagebox.showwarning(
                "Penanggung Jawab",
                "Pilih penanggung jawab kegiatan terlebih dahulu."
            )
            return

        self.penanggung_jawab_id = pj_id
        self.selected_ids.add(pj_id)

        conn = get_connection()
        cursor = conn.cursor()

        # Ambil anggota lama terlebih dahulu.
        # Ini penting supaya jadwal pegawai yang baru dihapus dari
        # kegiatan juga ikut dibersihkan otomatis.
        anggota_lama = {
            row[0]
            for row in cursor.execute(
                "SELECT pegawai_id FROM anggota_bagian WHERE bagian_id = ?",
                (self.bagian_id,)
            ).fetchall()
        }

        # Pegawai yang sebelumnya menjadi anggota tetapi sekarang
        # tidak dicentang lagi harus kehilangan seluruh jadwalnya
        # pada kegiatan ini.
        pegawai_dihapus = anggota_lama - self.selected_ids

        for pegawai_id in pegawai_dihapus:
            cursor.execute(
                """
                DELETE FROM jadwal
                WHERE bagian_id = ?
                AND pegawai_id = ?
                """,
                (self.bagian_id, pegawai_id)
            )

        nama_bagian = self.nama_bagian_var.get().strip() if hasattr(self, "nama_bagian_var") else ""
        if not nama_bagian:
            conn_nama = get_connection()
            row_nama = conn_nama.execute(
                "SELECT COALESCE(nama_bagian, '') FROM bagian WHERE id = ?",
                (self.bagian_id,)
            ).fetchone()
            conn_nama.close()
            nama_bagian = (row_nama[0] if row_nama else "").strip() or f"Kegiatan {self.bagian_aktif}"

        cursor.execute("""
            UPDATE bagian
            SET nama_bagian = ?,
                kegiatan = ?,
                lokasi = ?,
                dana = ?,
                penanggung_jawab_id = ?
            WHERE id = ?
        """, (
            nama_bagian,
            self.get_deskripsi(),
            self.lokasi_var.get().strip(),
            dana,
            pj_id,
            self.bagian_id
        ))

        cursor.execute(
            "DELETE FROM anggota_bagian WHERE bagian_id = ?",
            (self.bagian_id,)
        )

        for pegawai_id in self.selected_ids:
            cursor.execute("""
                INSERT INTO anggota_bagian
                (bagian_id, pegawai_id)
                VALUES (?, ?)
            """, (self.bagian_id, pegawai_id))

        conn.commit()
        conn.close()

        self.refresh_selection_table()
        self.refresh_bagian_buttons()
        self.load_kalender()
        self.refresh_data_kegiatan()
        self.refresh_data_pegawai()
        self.load_kalender_keseluruhan()

        nama_pj = next(
            (p[1] for p in self.pegawai if p[0] == pj_id),
            "-"
        )

        messagebox.showinfo(
            "Berhasil",
            f"Kegiatan {self.bagian_aktif} berhasil disimpan.\n\n"
            f"Penanggung jawab: {nama_pj}"
        )

    def bersihkan_bagian(self):

        if not self.bagian_id:
            return

        konfirmasi = messagebox.askyesno(
            "Hapus Data Kegiatan",
            f"Yakin ingin menghapus seluruh data "
            f"KEGIATAN {self.bagian_aktif}?\n\n"
            f"Yang akan dihapus:\n"
            f"• Pegawai kegiatan ini\n"
            f"• Jadwal penugasan kegiatan ini\n"
            f"• Kegiatan\n"
            f"• Lokasi\n"
            f"• Dana bulanan\n\n"
            f"Data kegiatan lain tidak akan terpengaruh."
        )

        if not konfirmasi:
            return

        conn = get_connection()
        cursor = conn.cursor()

        # Hapus jadwal kegiatan ini
        cursor.execute(
            """
            DELETE FROM jadwal
            WHERE bagian_id = ?
            """,
            (self.bagian_id,)
        )

        # Hapus daftar pegawai bagian
        cursor.execute(
            """
            DELETE FROM anggota_bagian
            WHERE bagian_id = ?
            """,
            (self.bagian_id,)
        )

        # Reset data bagian
        cursor.execute(
            """
            UPDATE bagian
            SET kegiatan = '',
                lokasi = '',
                dana = 0,
                penanggung_jawab_id = NULL
            WHERE id = ?
            """,
            (self.bagian_id,)
        )

        conn.commit()
        conn.close()

        # Reset tampilan
        self.set_deskripsi("")
        self.lokasi_var.set("")
        self.dana_var.set("")

        self.selected_ids = set()

        if hasattr(self, "komposisi_aktif"):
            del self.komposisi_aktif

        self.btn_acak_lagi.config(
            state="disabled"
        )

        self.refresh_selection_table()
        self.load_kalender()

        messagebox.showinfo(
            "Berhasil",
            f"Data KEGIATAN {self.bagian_aktif} "
            f"berhasil dibersihkan."
        )     
        
        # =====================================================
        # Acak Pegawai
        # =====================================================
    
    def acak_pegawai_dari_komposisi(self, hasil_item):
        """Acak pegawai sesuai komposisi, tetapi PJ selalu dipertahankan."""
        kelompok = {}
        for i, p in enumerate(self.pegawai):
            key = (p[5] or "Tidak diketahui", p[7])
            kelompok.setdefault(key, []).append(i)

        pegawai_terpilih = []
        pj_id = getattr(self, "penanggung_jawab_id", None)
        pj_index = next((i for i, p in enumerate(self.pegawai) if p[0] == pj_id), None)

        if pj_index is None:
            messagebox.showerror(
                "Penanggung Jawab",
                "Penanggung jawab belum dipilih."
            )
            return []

        pegawai_terpilih.append(pj_index)
        pj_key = (
            self.pegawai[pj_index][5] or "Tidak diketahui",
            self.pegawai[pj_index][7]
        )

        for golongan, tarif, jumlah in hasil_item["signature"]:
            key = (golongan, tarif)
            wajib = 1 if key == pj_key else 0
            jumlah_acak = jumlah - wajib

            if jumlah_acak <= 0:
                continue

            daftar_index = [
                i for i in kelompok.get(key, [])
                if i != pj_index
            ]

            if len(daftar_index) < jumlah_acak:
                messagebox.showerror(
                    "Data Tidak Cukup",
                    f"Pegawai {golongan} dengan tarif {rupiah(tarif)} tidak mencukupi."
                )
                return []

            riwayat = load_riwayat_turun(self.periode_id)
            random.shuffle(daftar_index)
            daftar_index.sort(
                key=lambda i: (
                    riwayat.get(self.pegawai[i][0], {}).get("bulan_ini", 0),
                    riwayat.get(self.pegawai[i][0], {}).get("total", 0)
                )
            )

            pegawai_terpilih.extend(daftar_index[:jumlah_acak])

        return pegawai_terpilih

    def acak_lagi(self):

        if not hasattr(self, "komposisi_aktif") or not self.komposisi_aktif:
            messagebox.showwarning(
                "Acak Pegawai",
                "Belum ada kombinasi yang dipilih. Gunakan Auto Tentukan terlebih dahulu."
            )
            return

        # Pastikan dana yang sekarang masih sama dengan target komposisi.
        try:
            dana_sekarang = parse_rupiah(self.dana_var.get())
        except Exception:
            messagebox.showwarning("Dana Tidak Valid", "Masukkan dana bulanan yang valid terlebih dahulu.")
            return

        target_komposisi = sum(
            int(tarif) * int(jumlah)
            for _gol, tarif, jumlah in self.komposisi_aktif.get("signature", [])
        )

        if target_komposisi != dana_sekarang:
            messagebox.showwarning(
                "Komposisi Tidak Sesuai",
                "Dana bulanan sudah berubah sehingga komposisi sebelumnya tidak berlaku.\n\n"
                "Gunakan Auto Tentukan lagi untuk membuat komposisi baru sesuai dana saat ini."
            )
            return

        kombinasi = self.acak_pegawai_dari_komposisi(
            self.komposisi_aktif
        )

        if not kombinasi:
            return

        # Acak Lagi cukup mengganti pegawai sesuai komposisi yang sama.
        # Tidak perlu menampilkan dialog hasil setiap kali ditekan.
        self.terapkan_kombinasi(
            kombinasi,
            self.kombinasi_nomor,
            tampilkan_info=False
        )

    def terapkan_kombinasi(self, kombinasi, nomor_kombinasi, tampilkan_info=True):
        # Simpan ID pegawai
        
        self.selected_ids = set(
            self.pegawai[i][0]
            for i in kombinasi
        )
    
        self.refresh_selection_table()
    
        # Hitung rincian golongan
        rincian_golongan = {}
    
        for i in kombinasi:
            pegawai_data = self.pegawai[i]
    
            golongan = (
                pegawai_data[5]
                or "Tidak diketahui"
            )
    
            tarif = pegawai_data[7]
    
            if golongan not in rincian_golongan:
                rincian_golongan[golongan] = {
                    "jumlah": 0,
                    "total": 0
                }
    
            rincian_golongan[golongan]["jumlah"] += 1
            rincian_golongan[golongan]["total"] += tarif
    
        total = sum(
            self.pegawai[i][7]
            for i in kombinasi
        )
    
        pesan_final = (
            f"Kegiatan {self.bagian_aktif}\n\n"
            f"KOMBINASI {nomor_kombinasi} DIPILIH\n"
            f"────────────────────\n"
        )
    
        for golongan in sorted(
            rincian_golongan.keys()
        ):
            data = rincian_golongan[golongan]
    
            pesan_final += (
                f"Gol {golongan:<6} : "
                f"{data['jumlah']} orang   "
                f"{rupiah(data['total'])}\n"
            )
    
        dana_target = parse_rupiah(self.dana_var.get())
        selisih = dana_target - total

        pesan_final += (
            f"\n────────────────────\n"
            f"Total Pegawai : {len(kombinasi)} orang\n"
            f"Total Dana    : {rupiah(total)}\n"
        )

        if selisih == 0:
            pesan_final += "Dana terpenuhi: Rp 0\n"
        elif selisih > 0:
            pesan_final += f"Dana belum terpakai: {rupiah(selisih)}\n"
        else:
            # Seharusnya tidak terjadi untuk kombinasi hasil Auto Tentukan.
            pesan_final += f"PERINGATAN: melebihi dana {rupiah(abs(selisih))}\n"

        if tampilkan_info:
            messagebox.showinfo(
                "Hasil Penentuan",
                pesan_final
            )

    # =====================================================
    # Hapus data bagian 
    # =====================================================

    def hapus_data_bagian(self, kode=None):
        """Hapus kegiatan tertentu. Jika kode=None, hapus kegiatan aktif."""
        kode_target = kode or self.bagian_aktif

        kode_list = load_bagian_codes(self.periode_id)
        if len(kode_list) <= 1:
            messagebox.showwarning(
                "Tidak Bisa Dihapus",
                "Minimal harus ada satu kegiatan dalam satu periode."
            )
            return

        data_target = load_bagian(self.periode_id, kode_target)
        if not data_target:
            self.refresh_bagian_buttons()
            return

        bagian_id_target = data_target[0]

        conn = get_connection()
        nama_row = conn.execute(
            "SELECT COALESCE(nama_bagian, '') FROM bagian WHERE id = ?",
            (bagian_id_target,)
        ).fetchone()
        conn.close()
        nama_target = (nama_row[0] if nama_row else "").strip() or f"KEGIATAN {kode_target}"

        konfirmasi = messagebox.askyesno(
            "Hapus Kegiatan",
            f'Yakin ingin menghapus "{nama_target}"?\n\n'
            "Kegiatan, pegawai, dan seluruh jadwal kegiatan ini akan dihapus."
        )
        if not konfirmasi:
            return

        conn = None
        try:
            conn = get_connection()
            cursor = conn.cursor()
            cursor.execute("DELETE FROM jadwal WHERE bagian_id = ?", (bagian_id_target,))
            cursor.execute("DELETE FROM anggota_bagian WHERE bagian_id = ?", (bagian_id_target,))
            cursor.execute("DELETE FROM bagian WHERE id = ?", (bagian_id_target,))
            conn.commit()
        except sqlite3.OperationalError as e:
            if conn:
                conn.rollback()
            if "locked" in str(e).lower():
                messagebox.showerror(
                    "Database Sedang Dipakai",
                    "Database sedang terkunci. Tutup program/Excel lain yang sedang membuka database, lalu coba lagi."
                )
            else:
                messagebox.showerror("Gagal Menghapus", str(e))
            return
        except Exception as e:
            if conn:
                conn.rollback()
            messagebox.showerror("Gagal Menghapus", str(e))
            return
        finally:
            if conn:
                conn.close()

        kode_list = load_bagian_codes(self.periode_id)
        if not kode_list:
            # Pengaman tambahan: periode tidak boleh kehilangan semua bagian.
            self.tambah_bagian()
            return

        if self.bagian_aktif == kode_target:
            self.bagian_aktif = kode_list[0]

        self.selected_ids = set()
        self.penanggung_jawab_id = None
        self.komposisi_aktif = None

        self.refresh_bagian_buttons()
        self.load_bagian_data()
        self.load_kalender()
        self.refresh_data_pegawai()
        self.refresh_selection_table()

        messagebox.showinfo("Berhasil", f'"{nama_target}" berhasil dihapus beserta seluruh jadwalnya.')

    # =====================================================
    # AUTO TENTUKAN
    # =====================================================

    def auto_tentukan(self):
        try:
            dana = parse_rupiah(self.dana_var.get())
        except Exception:
            messagebox.showerror("Dana Tidak Valid", "Masukkan dana bulanan yang valid.")
            return

        if dana <= 0:
            messagebox.showwarning("Dana Kosong", "Masukkan dana bulanan terlebih dahulu.")
            return

        pj_id = self.get_selected_pj_id()
        if not pj_id:
            messagebox.showwarning(
                "Penanggung Jawab",
                "Pilih penanggung jawab sebelum menggunakan Auto Tentukan."
            )
            return

        self.penanggung_jawab_id = pj_id
        pj = next((p for p in self.pegawai if p[0] == pj_id), None)
        if not pj:
            messagebox.showerror("Data Pegawai", "Data penanggung jawab tidak ditemukan.")
            return

        tarif_pj = pj[7]
        sisa_dana = dana - tarif_pj

        if sisa_dana < 0:
            messagebox.showwarning(
                "Dana Tidak Cukup",
                f"Tarif penanggung jawab {rupiah(tarif_pj)} lebih besar dari dana {rupiah(dana)}."
            )
            return

        # Cari kombinasi pegawai lain untuk sisa dana, tanpa PJ.
        kandidat = [p for p in self.pegawai if p[0] != pj_id]
        if sisa_dana == 0:
            hasil = [{
                "nomor": 1,
                "signature": [(pj[5] or "Tidak diketahui", pj[7], 1)]
            }]
        else:
            # Ambil seluruh komposisi yang bisa dibentuk, bukan hanya 5.
            hasil_lain = cari_beberapa_kombinasi(
                kandidat,
                sisa_dana,
                max_hasil=None
            )
            hasil = []
            pj_key = (pj[5] or "Tidak diketahui", pj[7])
            for item in hasil_lain:
                signature = list(item["signature"])
                # Tambahkan PJ ke kelompok golongan/tarif yang sama.
                found = False
                for n, (gol, tarif, jumlah) in enumerate(signature):
                    if (gol, tarif) == pj_key:
                        signature[n] = (gol, tarif, jumlah + 1)
                        found = True
                        break
                if not found:
                    signature.append((pj_key[0], pj_key[1], 1))
                signature.sort(key=lambda x: (x[0], x[1]))
                hasil.append({
                    "nomor": len(hasil) + 1,
                    "signature": signature
                })

        if not hasil:
            bawah, atas = cari_saran_dana_terdekat(kandidat, sisa_dana)

            kandidat_saran = []
            if bawah is not None:
                kandidat_saran.append((bawah + tarif_pj, abs(dana - (bawah + tarif_pj))))
            if atas is not None:
                kandidat_saran.append((atas + tarif_pj, abs(dana - (atas + tarif_pj))))

            kandidat_saran = sorted(set(kandidat_saran), key=lambda x: (x[1], x[0]))
            pesan_saran = (
                f"Tidak ditemukan kombinasi pegawai yang tepat untuk dana {rupiah(dana)}.\n\n"
                f"Dana yang harus dipenuhi setelah PJ: {rupiah(sisa_dana)}.\n"
            )

            if kandidat_saran:
                pesan_saran += "\nSaran dana terdekat:"
                for total_saran, selisih in kandidat_saran[:2]:
                    arah = "lebih rendah" if total_saran < dana else "lebih tinggi" if total_saran > dana else "sama"
                    pesan_saran += (
                        f"\n• {rupiah(total_saran)} ({arah}, selisih {rupiah(selisih)})"
                    )
            else:
                pesan_saran += "\nBelum ada kombinasi dana yang dapat dibentuk dari pegawai yang tersedia."

            messagebox.showwarning(
                "Tidak Ada Kombinasi",
                pesan_saran
            )
            return

        nama_pj = pj[1]
        pesan = (
            f"Kegiatan {self.bagian_aktif}\n"
            f"Dana: {rupiah(dana)}\n"
            f"PJ wajib: {nama_pj} ({rupiah(tarif_pj)})\n\n"
            f"ALTERNATIF KOMBINASI\n"
            f"────────────────────\n"
        )

        for item in hasil:
            pesan += f"\nKombinasi {item['nomor']}\n"
            rincian = {}
            for golongan, tarif, jumlah in item["signature"]:
                rincian.setdefault(golongan, {"jumlah": 0, "total": 0})
                rincian[golongan]["jumlah"] += jumlah
                rincian[golongan]["total"] += jumlah * tarif
            for golongan in sorted(rincian):
                data = rincian[golongan]
                pesan += f"Gol {golongan:<6}: {data['jumlah']} orang   {rupiah(data['total'])}\n"
            pesan += f"Total       : {sum(x['jumlah'] for x in rincian.values())} orang\n"

        pilihan = self.pilih_kombinasi_dialog(hasil, dana, nama_pj)

        if pilihan is None:
            return

        self.komposisi_aktif = hasil[pilihan - 1]
        self.kombinasi_nomor = pilihan

        kombinasi = self.acak_pegawai_dari_komposisi(self.komposisi_aktif)
        if not kombinasi:
            return

        if hasattr(self, "btn_acak_lagi"):
            self.btn_acak_lagi.config(state="normal")

        self.terapkan_kombinasi(kombinasi, pilihan)

    # =====================================================
    # KALENDER
    # =====================================================

    def label_kegiatan(self, kode):
        """Nama kegiatan yang ramah untuk kalender dan pesan bentrok."""
        conn = get_connection()
        row = conn.execute(
            "SELECT COALESCE(nama_bagian, '') FROM bagian WHERE periode_id = ? AND kode = ?",
            (self.periode_id, kode)
        ).fetchone()
        conn.close()

        nama = (row[0] if row else "").strip()
        default = f"Kegiatan {kode}"
        if not nama or nama == default:
            return f"Kegiatan {kode}"
        return f"Kegiatan {kode} — {nama}"

    def ringkasan_kode_kegiatan(self):
        conn = get_connection()
        rows = conn.execute(
            "SELECT kode, COALESCE(nama_bagian, '') FROM bagian WHERE periode_id = ? ORDER BY id",
            (self.periode_id,)
        ).fetchall()
        conn.close()

        bagian = []
        for kode, nama in rows:
            nama = (nama or '').strip()
            if nama and nama != f"Kegiatan {kode}":
                bagian.append(f"{kode} = {nama}")
            else:
                bagian.append(f"{kode} = Kegiatan {kode}")
        return "  •  ".join(bagian)

    def update_notebook_title(self):
        if not hasattr(self, "notebook"):
            return
        nama = self.nama_bagian_var.get().strip() if hasattr(self, "nama_bagian_var") else f"Kegiatan {self.bagian_aktif}"
        self.notebook.tab(
            self.tab_pengaturan,
            text=f"  Pengaturan Kegiatan {self.bagian_aktif}  "
        )
        self.notebook.tab(
            self.tab_jadwal,
            text="  Kalender Kegiatan  "
        )
        self.notebook.tab(
            self.tab_data_kegiatan,
            text=f"  Data {nama}  "
        )

    def load_kalender(self):
        """Menampilkan kalender untuk kegiatan aktif; hanya pegawai yang dipilih pada kegiatan ini."""
        for widget in self.calendar_frame.winfo_children():
            widget.destroy()

        bulan = NAMA_BULAN.index(self.bulan_var.get())
        tahun = int(self.tahun_var.get())
        jumlah_hari = calendar.monthrange(tahun, bulan)[1]
        hari_libur = load_hari_libur(self.periode_id)

        # Ambil semua kegiatan pada periode aktif dan buat warna yang konsisten.
        conn = get_connection()
        kegiatan_rows = conn.execute(
            """
            SELECT kode, COALESCE(nama_bagian, ''), COALESCE(kegiatan, '')
            FROM bagian
            WHERE periode_id = ?
            ORDER BY id
            """,
            (self.periode_id,)
        ).fetchall()

        jadwal_rows = conn.execute(
            """
            SELECT j.pegawai_id, j.hari, b.kode, COALESCE(b.nama_bagian, ''), COALESCE(b.kegiatan, '')
            FROM jadwal j
            JOIN bagian b ON b.id = j.bagian_id
            WHERE j.periode_id = ?
            """,
            (self.periode_id,)
        ).fetchall()
        conn.close()

        palette = [
            "#B8D4EF", "#C7E6C5", "#F3D6A5", "#D9C2E9",
            "#F4B8B8", "#BFE3E3", "#E7D6A8", "#C9D7B8",
            "#D5C4A1", "#C7D4F2", "#E5C1CD", "#CDE4C8"
        ]
        self.calendar_colors = {
            row[0]: palette[i % len(palette)]
            for i, row in enumerate(kegiatan_rows)
        }
        self.calendar_activity_names = {}
        for kode, nama, deskripsi in kegiatan_rows:
            nama_tampil = (nama or '').strip() or f"Kegiatan {kode}"
            self.calendar_activity_names[kode] = nama_tampil

        jadwal_map = {
            (row[0], row[1]): {
                "kode": row[2],
                "nama": (row[3] or '').strip() or f"Kegiatan {row[2]}",
                "deskripsi": (row[4] or '').strip(),
            }
            for row in jadwal_rows
        }

        if hasattr(self, "info_kalender"):
            self.info_kalender.config(
                text=(
                    f"Pegawai terpilih pada kegiatan ini • {NAMA_BULAN[bulan]} {tahun}  •  "
                    f"Kegiatan aktif: {self.bagian_aktif}"
                )
            )

        if hasattr(self, "info_libur"):
            if hari_libur:
                self.info_libur.config(
                    text=f"Libur manual: {', '.join(str(x) for x in sorted(hari_libur))}  •  Minggu selalu libur"
                )
            else:
                self.info_libur.config(
                    text="Libur manual: tidak ada  •  Minggu selalu libur"
                )

        # Legenda kode kegiatan + warna.
        if hasattr(self, "legend_kegiatan"):
            for widget in self.legend_kegiatan.winfo_children():
                widget.destroy()

            tk.Label(
                self.legend_kegiatan,
                text="Keterangan kegiatan:",
                bg="#F4F6F8",
                font=("Segoe UI", 9, "bold")
            ).pack(side="left", padx=(0, 8))

            for kode, nama, deskripsi in kegiatan_rows:
                warna = self.calendar_colors.get(kode, "#DDE7F0")
                label = tk.Label(
                    self.legend_kegiatan,
                    text=f"{kode} = {(nama or '').strip() or f'Kegiatan {kode}'}",
                    bg=warna,
                    relief="solid",
                    bd=1,
                    padx=8,
                    pady=3,
                    font=("Segoe UI", 9, "bold")
                )
                label.pack(side="left", padx=3)

        nama_hari = ["Sen", "Sel", "Rab", "Kam", "Jum", "Sab", "Min"]

        ttk.Label(
            self.calendar_frame,
            text="Nama Pegawai",
            font=("Segoe UI", 9, "bold"),
            anchor="center"
        ).grid(row=0, column=0, rowspan=2, sticky="nsew", padx=1, pady=1)

        for hari in range(1, jumlah_hari + 1):
            tanggal = datetime(tahun, bulan, hari)
            index_hari = tanggal.weekday()
            nonkerja = index_hari == 6 or hari in hari_libur

            label_tanggal = ttk.Label(
                self.calendar_frame,
                text=str(hari),
                font=("Segoe UI", 9, "bold"),
                anchor="center"
            )
            label_tanggal.grid(row=0, column=hari, sticky="nsew", padx=1, pady=(1, 0))

            label_hari = ttk.Label(
                self.calendar_frame,
                text=("LIBUR" if nonkerja else nama_hari[index_hari]),
                font=("Segoe UI", 7 if nonkerja else 8, "bold" if nonkerja else "normal"),
                anchor="center"
            )
            label_hari.grid(row=1, column=hari, sticky="nsew", padx=1, pady=(0, 1))

            if nonkerja:
                label_tanggal.configure(style="Holiday.TLabel")
                label_hari.configure(style="Holiday.TLabel")

        # Kalender kegiatan hanya menampilkan pegawai yang dicentang/terdaftar
        # pada kegiatan yang sedang aktif. Kalender keseluruhan memakai semua pegawai.
        conn = get_connection()
        anggota_rows = conn.execute(
            "SELECT pegawai_id FROM anggota_bagian WHERE bagian_id = ?",
            (self.bagian_id,)
        ).fetchall() if self.bagian_id else []
        conn.close()
        anggota_ids = {row[0] for row in anggota_rows}
        semua_pegawai = sorted(
            [p for p in self.pegawai if p[0] in anggota_ids],
            key=lambda p: str(p[1] or '').lower()
        )

        if not semua_pegawai:
            ttk.Label(
                self.calendar_frame,
                text="Belum ada pegawai yang dipilih untuk kegiatan ini.",
                font=("Segoe UI", 10)
            ).grid(row=2, column=0, columnspan=jumlah_hari + 1, pady=20)
            return

        def show_detail(pegawai_nama, hari, detail):
            if not hasattr(self, "calendar_detail"):
                return
            if detail:
                self.calendar_detail.config(
                    text=f"{pegawai_nama} • Tanggal {hari} • {detail['kode']} — {detail['nama']}"
                )
            else:
                self.calendar_detail.config(
                    text=f"{pegawai_nama} • Tanggal {hari} • Belum ada penugasan"
                )

        for row, pegawai in enumerate(semua_pegawai, start=2):
            pegawai_id = pegawai[0]
            nama = pegawai[1] or "Tanpa Nama"

            ttk.Label(
                self.calendar_frame,
                text=nama,
                anchor="w",
                padding=(5, 0)
            ).grid(row=row, column=0, sticky="nsew", padx=1, pady=1)

            for hari in range(1, jumlah_hari + 1):
                index_hari = calendar.weekday(tahun, bulan, hari)
                nonkerja = index_hari == 6 or hari in hari_libur
                detail = jadwal_map.get((pegawai_id, hari))

                if nonkerja:
                    tombol = tk.Button(
                        self.calendar_frame,
                        text="LIBUR",
                        width=6,
                        height=1,
                        relief="solid",
                        borderwidth=1,
                        state="disabled",
                        disabledforeground="#888888",
                        bg="#ECECEC"
                    )
                else:
                    kode = detail["kode"] if detail else ""
                    warna = self.calendar_colors.get(kode, "#FFFFFF") if detail else "#FFFFFF"
                    tombol = tk.Button(
                        self.calendar_frame,
                        text=kode,
                        width=6,
                        height=1,
                        relief="solid",
                        borderwidth=1,
                        bg=warna,
                        activebackground=warna,
                        font=("Segoe UI", 9, "bold") if kode else ("Segoe UI", 9),
                        command=lambda pid=pegawai_id, h=hari: self.toggle_jadwal(pid, h)
                    )

                tombol.bind(
                    "<Enter>",
                    lambda event, n=nama, h=hari, d=detail: show_detail(n, h, d)
                )
                tombol.bind(
                    "<Leave>",
                    lambda event: self.calendar_detail.config(text="") if hasattr(self, "calendar_detail") else None
                )
                tombol.grid(row=row, column=hari, sticky="nsew", padx=1, pady=1)

        self.calendar_frame.grid_columnconfigure(0, minsize=220)
        for col in range(1, jumlah_hari + 1):
            self.calendar_frame.grid_columnconfigure(col, minsize=48)

        self.load_kalender_keseluruhan()

    def load_kalender_keseluruhan(self):
        """Kalender dashboard bulanan: semua pegawai dan semua kegiatan, bersifat baca-saja."""
        if (
            getattr(self, "overall_calendar_window", None) is None
            or not self.overall_calendar_window.winfo_exists()
            or not hasattr(self, "overall_calendar_frame")
        ):
            return

        for widget in self.overall_calendar_frame.winfo_children():
            widget.destroy()

        bulan = NAMA_BULAN.index(self.bulan_var.get())
        tahun = int(self.tahun_var.get())
        jumlah_hari = calendar.monthrange(tahun, bulan)[1]
        hari_libur = load_hari_libur(self.periode_id)

        conn = get_connection()
        kegiatan_rows = conn.execute(
            """
            SELECT kode, COALESCE(nama_bagian, ''), COALESCE(kegiatan, '')
            FROM bagian
            WHERE periode_id = ?
            ORDER BY id
            """,
            (self.periode_id,)
        ).fetchall()

        jadwal_rows = conn.execute(
            """
            SELECT j.pegawai_id, j.hari, b.kode, COALESCE(b.nama_bagian, ''),
                   COALESCE(b.kegiatan, '')
            FROM jadwal j
            JOIN bagian b ON b.id = j.bagian_id
            WHERE j.periode_id = ?
            """,
            (self.periode_id,)
        ).fetchall()
        conn.close()

        palette = [
            "#B8D4EF", "#C7E6C5", "#F3D6A5", "#D9C2E9",
            "#F4B8B8", "#BFE3E3", "#E7D6A8", "#C9D7B8",
            "#D5C4A1", "#C7D4F2", "#E5C1CD", "#CDE4C8"
        ]
        colors = {row[0]: palette[i % len(palette)] for i, row in enumerate(kegiatan_rows)}
        names = {row[0]: ((row[1] or '').strip() or f"Kegiatan {row[0]}") for row in kegiatan_rows}
        jadwal_map = {
            (row[0], row[1]): {
                "kode": row[2],
                "nama": (row[3] or '').strip() or f"Kegiatan {row[2]}",
                "deskripsi": (row[4] or '').strip(),
            }
            for row in jadwal_rows
        }

        self.overall_calendar_detail.config(text="")
        self.overall_info_kalender.config(
            text=f"Semua pegawai • {NAMA_BULAN[bulan]} {tahun} • Semua kegiatan"
        )
        if hari_libur:
            self.overall_info_libur.config(
                text=f"Libur manual: {', '.join(str(x) for x in sorted(hari_libur))} • Minggu selalu libur"
            )
        else:
            self.overall_info_libur.config(text="Libur manual: tidak ada • Minggu selalu libur")

        for widget in self.overall_legend_kegiatan.winfo_children():
            widget.destroy()
        tk.Label(
            self.overall_legend_kegiatan,
            text="Keterangan kegiatan:",
            bg="#F4F6F8",
            font=("Segoe UI", 9, "bold")
        ).pack(side="left", padx=(0, 8))
        for kode, nama, deskripsi in kegiatan_rows:
            tk.Label(
                self.overall_legend_kegiatan,
                text=f"{kode} = {names[kode]}",
                bg=colors[kode],
                relief="solid", bd=1, padx=8, pady=3,
                font=("Segoe UI", 9, "bold")
            ).pack(side="left", padx=3)

        nama_hari = ["Sen", "Sel", "Rab", "Kam", "Jum", "Sab", "Min"]
        ttk.Label(
            self.overall_calendar_frame, text="Nama Pegawai",
            font=("Segoe UI", 9, "bold"), anchor="center"
        ).grid(row=0, column=0, rowspan=2, sticky="nsew", padx=1, pady=1)

        for hari in range(1, jumlah_hari + 1):
            tanggal = datetime(tahun, bulan, hari)
            nonkerja = tanggal.weekday() == 6 or hari in hari_libur
            ttk.Label(
                self.overall_calendar_frame, text=str(hari),
                font=("Segoe UI", 9, "bold"), anchor="center",
                style="Holiday.TLabel" if nonkerja else "TLabel"
            ).grid(row=0, column=hari, sticky="nsew", padx=1, pady=(1, 0))
            ttk.Label(
                self.overall_calendar_frame,
                text="LIBUR" if nonkerja else nama_hari[tanggal.weekday()],
                font=("Segoe UI", 7 if nonkerja else 8, "bold" if nonkerja else "normal"),
                anchor="center",
                style="Holiday.TLabel" if nonkerja else "TLabel"
            ).grid(row=1, column=hari, sticky="nsew", padx=1, pady=(0, 1))

        semua_pegawai = sorted(self.pegawai, key=lambda p: str(p[1] or '').lower())
        if not semua_pegawai:
            ttk.Label(
                self.overall_calendar_frame, text="Belum ada data pegawai aktif.",
                font=("Segoe UI", 10)
            ).grid(row=2, column=0, columnspan=jumlah_hari + 1, pady=20)
            return

        for row, pegawai in enumerate(semua_pegawai, start=2):
            pegawai_id = pegawai[0]
            nama = pegawai[1] or "Tanpa Nama"
            ttk.Label(
                self.overall_calendar_frame, text=nama, anchor="w", padding=(5, 0)
            ).grid(row=row, column=0, sticky="nsew", padx=1, pady=1)

            for hari in range(1, jumlah_hari + 1):
                tanggal = datetime(tahun, bulan, hari)
                nonkerja = tanggal.weekday() == 6 or hari in hari_libur
                detail = jadwal_map.get((pegawai_id, hari))
                kode = detail["kode"] if detail else ""
                warna = colors.get(kode, "#FFFFFF") if detail else "#FFFFFF"

                if nonkerja:
                    cell = tk.Label(
                        self.overall_calendar_frame, text="LIBUR", width=6, height=1,
                        relief="solid", bd=1, bg="#ECECEC", fg="#888888",
                        font=("Segoe UI", 8)
                    )
                else:
                    cell = tk.Label(
                        self.overall_calendar_frame, text=kode, width=6, height=1,
                        relief="solid", bd=1, bg=warna,
                        font=("Segoe UI", 9, "bold") if kode else ("Segoe UI", 9)
                    )

                if detail:
                    detail_text = f"{nama} • Tanggal {hari} • {kode} — {detail['nama']}"
                    if detail.get("deskripsi"):
                        detail_text += f" • {detail['deskripsi']}"
                    cell.bind(
                        "<Enter>",
                        lambda event, txt=detail_text: self.overall_calendar_detail.config(text=txt)
                    )
                else:
                    cell.bind(
                        "<Enter>",
                        lambda event, n=nama, h=hari: self.overall_calendar_detail.config(
                            text=f"{n} • Tanggal {h} • Belum ada penugasan"
                        )
                    )
                cell.bind("<Leave>", lambda event: self.overall_calendar_detail.config(text=""))
                cell.grid(row=row, column=hari, sticky="nsew", padx=1, pady=1)

        self.overall_calendar_frame.grid_columnconfigure(0, minsize=220)
        for col in range(1, jumlah_hari + 1):
            self.overall_calendar_frame.grid_columnconfigure(col, minsize=48)

    def atur_hari_libur(self):
        """Dialog manual untuk menentukan tanggal merah pada bulan aktif."""
        bulan = NAMA_BULAN.index(self.bulan_var.get())
        tahun = int(self.tahun_var.get())
        jumlah_hari = calendar.monthrange(tahun, bulan)[1]
        tersimpan = load_hari_libur(self.periode_id)

        dialog = tk.Toplevel(self)
        dialog.title("Atur Hari Libur")
        dialog.transient(self)
        dialog.grab_set()
        dialog.resizable(False, False)

        tk.Label(
            dialog,
            text=f"Hari libur manual — {NAMA_BULAN[bulan]} {tahun}",
            font=("Segoe UI", 12, "bold")
        ).pack(anchor="w", padx=18, pady=(15, 3))

        tk.Label(
            dialog,
            text="Centang tanggal yang tidak boleh diisi. Hari Minggu otomatis libur.",
            fg="#666666"
        ).pack(anchor="w", padx=18, pady=(0, 10))

        grid = tk.Frame(dialog)
        grid.pack(padx=18, pady=5)

        vars_hari = {}
        nama_hari = ["Sen", "Sel", "Rab", "Kam", "Jum", "Sab", "Min"]

        for col, nama in enumerate(nama_hari):
            tk.Label(
                grid, text=nama, width=9, font=("Segoe UI", 9, "bold")
            ).grid(row=0, column=col, pady=(0, 5))

        for hari in range(1, jumlah_hari + 1):
            tanggal = datetime(tahun, bulan, hari)
            weekday = tanggal.weekday()
            row = 1 + (hari - 1) // 7
            col = weekday

            var = tk.BooleanVar(value=hari in tersimpan)
            vars_hari[hari] = var

            cb = tk.Checkbutton(
                grid,
                text=str(hari),
                variable=var,
                width=6,
                anchor="center"
            )
            cb.grid(row=row, column=col, padx=2, pady=2)

            if weekday == 6:
                cb.configure(
                    state="disabled",
                    text=f"{hari}*",
                    disabledforeground="#888888"
                )

        tk.Label(
            dialog,
            text="* Minggu selalu libur dan tidak perlu dicentang.",
            fg="#777777",
            font=("Segoe UI", 9, "italic")
        ).pack(anchor="w", padx=18, pady=(4, 8))

        tombol = tk.Frame(dialog)
        tombol.pack(fill="x", padx=18, pady=(4, 15))

        def simpan():
            baru = {
                hari for hari, var in vars_hari.items()
                if var.get() and datetime(tahun, bulan, hari).weekday() != 6
            }

            try:
                simpan_hari_libur(self.periode_id, baru)

                # Hapus jadwal yang jatuh pada Minggu atau libur manual
                # supaya kalender tidak menyimpan penugasan pada hari nonkerja.
                nonkerja = set(baru)
                nonkerja.update(
                    hari for hari in range(1, jumlah_hari + 1)
                    if datetime(tahun, bulan, hari).weekday() == 6
                )

                conn = get_connection()
                for hari in sorted(nonkerja):
                    conn.execute(
                        "DELETE FROM jadwal WHERE periode_id = ? AND hari = ?",
                        (self.periode_id, hari)
                    )
                conn.commit()
                conn.close()
            except Exception as e:
                try:
                    conn.close()
                except Exception:
                    pass
                messagebox.showerror("Gagal Menyimpan Hari Libur", str(e), parent=dialog)
                return

            dialog.destroy()
            self.load_kalender()
            self.refresh_data_pegawai()
            messagebox.showinfo(
                "Hari Libur",
                "Pengaturan hari libur bulan ini berhasil disimpan.",
                parent=self
            )

        ttk.Button(tombol, text="Batal", command=dialog.destroy).pack(side="right", padx=(5, 0))
        ttk.Button(tombol, text="Simpan", command=simpan).pack(side="right")

    # =====================================================
    # CELL KALENDER
    # =====================================================

    def create_calendar_cell(
        self,
        row,
        hari,
        pegawai_id
    ):

        # Cek jadwal
        conn = get_connection()
        cursor = conn.cursor()

        jadwal = cursor.execute("""
            SELECT
                j.bagian_id,
                b.kode
            FROM jadwal j
            JOIN bagian b
                ON b.id = j.bagian_id
            WHERE j.periode_id = ?
            AND j.pegawai_id = ?
            AND j.hari = ?
        """, (
            self.periode_id,
            pegawai_id,
            hari
        )).fetchone()

        conn.close()

        kode = jadwal[1] if jadwal else ""

        warna = "#DDE7F0"

        if kode == "A":
            warna = "#B8D4EF"

        elif kode == "B":
            warna = "#C7E6C5"

        elif kode == "C":
            warna = "#F3D6A5"

        button = tk.Button(
            self.calendar_frame,
            text=kode,
            width=3,
            height=1,
            bg=warna,
            relief="solid",
            bd=1,
            command=lambda
                p=pegawai_id,
                h=hari:
                self.toggle_jadwal(p, h)
        )

        button.grid(
            row=row,
            column=hari,
            padx=1,
            pady=1
        )

    # =====================================================
    # TOGGLE JADWAL
    # =====================================================

    def toggle_jadwal(
        self,
        pegawai_id,
        hari
    ):

        # Hari Minggu dan tanggal merah manual tidak dapat diisi.
        bulan = NAMA_BULAN.index(self.bulan_var.get())
        tahun = int(self.tahun_var.get())
        if datetime(tahun, bulan, hari).weekday() == 6 or hari in load_hari_libur(self.periode_id):
            return

        conn = get_connection()
        cursor = conn.cursor()

        existing = cursor.execute("""
            SELECT
                id,
                bagian_id
            FROM jadwal
            WHERE periode_id = ?
            AND pegawai_id = ?
            AND hari = ?
        """, (
            self.periode_id,
            pegawai_id,
            hari
        )).fetchone()

        if existing:

            if existing[1] == self.bagian_id:

                # Hapus jadwal sendiri
                cursor.execute("""
                    DELETE FROM jadwal
                    WHERE id = ?
                """, (
                    existing[0],
                ))

                conn.commit()
                conn.close()

                self.load_kalender()
                self.refresh_data_pegawai()

                return

            else:

                # Bentrok
                bagian_lain = cursor.execute("""
                    SELECT kode, COALESCE(nama_bagian, '')
                    FROM bagian
                    WHERE id = ?
                """, (
                    existing[1],
                )).fetchone()

                conn.close()

                kode_lain = bagian_lain[0] if bagian_lain else "?"
                nama_lain = (bagian_lain[1] if bagian_lain else "").strip() or f"Kegiatan {kode_lain}"
                nama_kegiatan_lain = f"Kegiatan {kode_lain} — {nama_lain}"
                messagebox.showwarning(
                    "Jadwal Bentrok",
                    f"Pegawai ini sudah bertugas pada {nama_kegiatan_lain} "
                    f"pada tanggal {hari}.\n\n"
                    f"Pegawai tidak dapat ditugaskan pada dua kegiatan "
                    f"di tanggal yang sama."
                )

                return

        # Tambah jadwal
        cursor.execute("""
            INSERT INTO jadwal
            (
                periode_id,
                bagian_id,
                pegawai_id,
                hari
            )
            VALUES (?, ?, ?, ?)
        """, (
            self.periode_id,
            self.bagian_id,
            pegawai_id,
            hari
        ))

        conn.commit()
        conn.close()

        self.load_kalender()

    def build_data_kegiatan(self):
        """Tabel rekap penugasan khusus kegiatan aktif pada bulan aktif."""
        for w in self.tab_data_kegiatan.winfo_children():
            w.destroy()

        tk.Label(
            self.tab_data_kegiatan, text="DATA KEGIATAN",
            bg="#F4F6F8", font=("Segoe UI", 16, "bold")
        ).pack(anchor="w", padx=20, pady=(15, 3))

        self.data_kegiatan_info = tk.Label(
            self.tab_data_kegiatan, text="", bg="#F4F6F8", fg="#555555"
        )
        self.data_kegiatan_info.pack(anchor="w", padx=20, pady=(0, 8))

        frame = tk.Frame(self.tab_data_kegiatan, bg="white")
        frame.pack(fill="both", expand=True, padx=20, pady=10)
        # Data Kegiatan hanya menampilkan JUMLAH TURUN pada kegiatan ini.
        # Tanggal/hari penugasan sudah ditampilkan pada Kalender Kegiatan,
        # sehingga tidak perlu diduplikasi di tabel data.
        cols = ("nama", "nip", "pangkat", "golongan", "jabatan", "tarif", "pj", "jumlah")
        self.table_data_kegiatan = ttk.Treeview(frame, columns=cols, show="headings", selectmode="browse")
        heads = {
            "nama":"Nama", "nip":"NIP", "pangkat":"Pangkat", "golongan":"Golongan",
            "jabatan":"Jabatan", "tarif":"Tarif", "pj":"PJ",
            "jumlah":"Turun"
        }
        widths = {"nama":280,"nip":185,"pangkat":170,"golongan":105,"jabatan":230,"tarif":125,"pj":70,"jumlah":120}
        for c in cols:
            self.table_data_kegiatan.heading(c, text=heads[c])
            self.table_data_kegiatan.column(c, width=widths[c], anchor="center")
        self.table_data_kegiatan.column("nama", anchor="w")
        self.table_data_kegiatan.column("jabatan", anchor="w")
        self.table_data_kegiatan.pack(side="left", fill="both", expand=True)
        scroll=ttk.Scrollbar(frame, orient="vertical", command=self.table_data_kegiatan.yview)
        scroll.pack(side="right", fill="y")
        self.table_data_kegiatan.configure(yscrollcommand=scroll.set)

        toolbar = tk.Frame(self.tab_data_kegiatan, bg="#F4F6F8")
        toolbar.pack(fill="x", padx=20, pady=(0, 12))
        ttk.Button(
            toolbar,
            text="🗑 Hapus Keanggotaan Permanen",
            command=self.hapus_keanggotaan_kegiatan_terpilih
        ).pack(side="left")
        tk.Label(
            toolbar,
            text="Hapus pegawai dari kegiatan ini pada bulan aktif beserta jadwalnya.",
            bg="#F4F6F8", fg="#666666", font=("Segoe UI", 9)
        ).pack(side="left", padx=10)
        self.table_data_kegiatan.bind(
            "<Delete>", lambda event: self.hapus_keanggotaan_kegiatan_terpilih()
        )
        self.refresh_data_kegiatan()

    def hapus_keanggotaan_kegiatan_terpilih(self):
        """Menghapus keanggotaan pegawai secara permanen dari kegiatan/periode aktif."""
        if not hasattr(self, "table_data_kegiatan") or self.table_data_kegiatan is None:
            return
        pilihan = self.table_data_kegiatan.selection()
        if not pilihan:
            messagebox.showinfo("Keanggotaan Kegiatan", "Pilih pegawai yang ingin dihapus dari kegiatan terlebih dahulu.", parent=self)
            return
        try:
            pegawai_id = int(pilihan[0])
        except (ValueError, TypeError):
            return

        if pegawai_id == getattr(self, "penanggung_jawab_id", None):
            messagebox.showwarning(
                "Penanggung Jawab",
                "Pegawai ini adalah Penanggung Jawab kegiatan.\n\nGanti Penanggung Jawab terlebih dahulu sebelum menghapus keanggotaannya.",
                parent=self
            )
            return

        conn = get_connection()
        try:
            row = conn.execute("SELECT nama FROM pegawai WHERE id = ?", (pegawai_id,)).fetchone()
            if not row:
                return
            nama = row[0] or "Tanpa Nama"
            jumlah_jadwal = conn.execute(
                "SELECT COUNT(*) FROM jadwal WHERE periode_id = ? AND bagian_id = ? AND pegawai_id = ?",
                (self.periode_id, self.bagian_id, pegawai_id)
            ).fetchone()[0]
        finally:
            conn.close()

        if not messagebox.askyesno(
            "Hapus Keanggotaan Permanen",
            f"Hapus {nama} dari Kegiatan {self.bagian_aktif} pada {NAMA_BULAN[self.bulan]} {self.tahun}?\n\n"
            f"• Keanggotaan kegiatan dihapus\n• {jumlah_jadwal} jadwal kegiatan ikut dihapus\n\n"
            "Pegawai tetap ada di Data Pegawai dan bisa ditambahkan kembali kapan saja.\n\nLanjutkan?",
            parent=self
        ):
            return

        conn = get_connection()
        try:
            conn.execute(
                "DELETE FROM jadwal WHERE periode_id = ? AND bagian_id = ? AND pegawai_id = ?",
                (self.periode_id, self.bagian_id, pegawai_id)
            )
            conn.execute(
                "DELETE FROM anggota_bagian WHERE bagian_id = ? AND pegawai_id = ?",
                (self.bagian_id, pegawai_id)
            )
            conn.commit()
        except Exception as exc:
            conn.rollback()
            messagebox.showerror("Gagal Menghapus", f"Keanggotaan gagal dihapus.\n\n{exc}", parent=self)
            return
        finally:
            conn.close()

        self.selected_ids.discard(pegawai_id)
        self.refresh_selection_table()
        self.refresh_data_kegiatan()
        self.refresh_data_pegawai()
        self.load_kalender()
        self.load_kalender_keseluruhan()

    def refresh_data_kegiatan(self):
        if not hasattr(self, "table_data_kegiatan") or self.table_data_kegiatan is None:
            return
        try:
            for item in self.table_data_kegiatan.get_children():
                self.table_data_kegiatan.delete(item)
        except tk.TclError:
            return

        data = load_bagian(self.periode_id, self.bagian_aktif)
        if not data:
            return
        bagian_id, kode, kegiatan, lokasi, dana, pj_id = data
        nama = getattr(self, "nama_bagian_var", tk.StringVar(value=f"Kegiatan {kode}")).get().strip() or f"Kegiatan {kode}"
        if hasattr(self, "data_kegiatan_info"):
            self.data_kegiatan_info.config(
                text=f"{nama} • {NAMA_BULAN[self.bulan]} {self.tahun} • hanya menghitung penugasan pada kegiatan ini"
            )

        conn=get_connection()
        try:
            anggota=conn.execute("""
                SELECT p.id,p.nama,p.nip,p.pangkat,p.golongan_asli,p.jabatan,p.tarif
                FROM anggota_bagian ab JOIN pegawai p ON p.id=ab.pegawai_id
                WHERE ab.bagian_id=? ORDER BY p.nama
            """, (bagian_id,)).fetchall()
            for p in anggota:
                # Satu-satunya angka rekap di Data Kegiatan adalah
                # jumlah turun pegawai pada kegiatan ini di bulan aktif.
                jumlah_turun = conn.execute(
                    "SELECT COUNT(*) FROM jadwal WHERE periode_id=? AND bagian_id=? AND pegawai_id=?",
                    (self.periode_id, bagian_id, p[0])
                ).fetchone()[0]
                self.table_data_kegiatan.insert(
                    "", "end", iid=str(p[0]),
                    values=(
                        p[1],p[2],p[3],p[4],p[5],rupiah(p[6]),
                        "Ya" if p[0]==pj_id else "",
                        f"{jumlah_turun} kali"
                    )
                )
        finally:
            conn.close()

    def build_jadwal(self):

        style = ttk.Style(self)
        style.configure(
            "Holiday.TLabel",
            foreground="#888888",
            background="#ECECEC"
        )

        tk.Label(
            self.tab_jadwal,
            text="KALENDER PENUGASAN KEGIATAN",
            bg="#F4F6F8",
            font=("Segoe UI", 16, "bold")
        ).pack(
            anchor="w",
            padx=15,
            pady=(15, 3)
        )

        self.info_kalender = tk.Label(
            self.tab_jadwal,
            text="",
            bg="#F4F6F8",
            fg="#555555"
        )

        self.info_kalender.pack(
            anchor="w",
            padx=15
        )

        kalender_tools = tk.Frame(
            self.tab_jadwal,
            bg="#F4F6F8"
        )
        kalender_tools.pack(
            fill="x",
            padx=15,
            pady=(8, 0)
        )

        ttk.Button(
            kalender_tools,
            text="⚙ Atur Hari Libur",
            command=self.atur_hari_libur
        ).pack(side="left")

        self.info_libur = tk.Label(
            kalender_tools,
            text="",
            bg="#F4F6F8",
            fg="#666666"
        )
        self.info_libur.pack(side="left", padx=12)

        self.legend_kegiatan = tk.Frame(
            self.tab_jadwal,
            bg="#F4F6F8"
        )
        self.legend_kegiatan.pack(
            fill="x",
            padx=15,
            pady=(8, 0)
        )

        self.calendar_detail = tk.Label(
            self.tab_jadwal,
            text="",
            bg="#F4F6F8",
            fg="#555555",
            anchor="w"
        )
        self.calendar_detail.pack(
            fill="x",
            padx=15,
            pady=(4, 0)
        )

        outer = tk.Frame(
            self.tab_jadwal
        )

        outer.pack(
            fill="both",
            expand=True,
            padx=15,
            pady=10
        )

        self.calendar_canvas = tk.Canvas(
            outer,
            bg="white"
        )

        hscroll = ttk.Scrollbar(
            outer,
            orient="horizontal",
            command=self.calendar_canvas.xview
        )

        vscroll = ttk.Scrollbar(
            outer,
            orient="vertical",
            command=self.calendar_canvas.yview
        )

        self.calendar_canvas.configure(
            xscrollcommand=hscroll.set,
            yscrollcommand=vscroll.set
        )

        self.calendar_canvas.grid(
            row=0,
            column=0,
            sticky="nsew"
        )

        vscroll.grid(
            row=0,
            column=1,
            sticky="ns"
        )

        hscroll.grid(
            row=1,
            column=0,
            sticky="ew"
        )

        outer.rowconfigure(
            0,
            weight=1
        )

        outer.columnconfigure(
            0,
            weight=1
        )

        self.calendar_frame = tk.Frame(
            self.calendar_canvas,
            bg="white"
        )

        self.calendar_window = (
            self.calendar_canvas.create_window(
                (0, 0),
                window=self.calendar_frame,
                anchor="nw"
            )
        )

        self.calendar_frame.bind(
            "<Configure>",
            lambda e: self.calendar_canvas.configure(
                scrollregion=self.calendar_canvas.bbox("all")
            )
        )

    def build_kalender_keseluruhan(self, parent):
        """Membangun jendela kalender keseluruhan untuk semua pegawai dan kegiatan."""
        ttk.Style(self).configure(
            "Holiday.TLabel", foreground="#888888", background="#ECECEC"
        )

        tk.Label(
            parent,
            text="KALENDER KESELURUHAN",
            bg="#F4F6F8", font=("Segoe UI", 16, "bold")
        ).pack(anchor="w", padx=15, pady=(15, 3))

        self.overall_info_kalender = tk.Label(
            parent, text="",
            bg="#F4F6F8", fg="#555555"
        )
        self.overall_info_kalender.pack(anchor="w", padx=15)

        tools = tk.Frame(parent, bg="#F4F6F8")
        tools.pack(fill="x", padx=15, pady=(8, 0))
        ttk.Button(
            tools, text="⚙ Atur Hari Libur", command=self.atur_hari_libur
        ).pack(side="left")
        self.overall_info_libur = tk.Label(
            tools, text="", bg="#F4F6F8", fg="#666666"
        )
        self.overall_info_libur.pack(side="left", padx=12)

        self.overall_legend_kegiatan = tk.Frame(
            parent, bg="#F4F6F8"
        )
        self.overall_legend_kegiatan.pack(fill="x", padx=15, pady=(8, 0))

        self.overall_calendar_detail = tk.Label(
            parent, text="", bg="#F4F6F8",
            fg="#555555", anchor="w"
        )
        self.overall_calendar_detail.pack(fill="x", padx=15, pady=(4, 0))

        outer = tk.Frame(parent)
        outer.pack(fill="both", expand=True, padx=15, pady=10)

        canvas = tk.Canvas(outer, bg="white")
        hscroll = ttk.Scrollbar(outer, orient="horizontal", command=canvas.xview)
        vscroll = ttk.Scrollbar(outer, orient="vertical", command=canvas.yview)
        canvas.configure(xscrollcommand=hscroll.set, yscrollcommand=vscroll.set)
        canvas.grid(row=0, column=0, sticky="nsew")
        vscroll.grid(row=0, column=1, sticky="ns")
        hscroll.grid(row=1, column=0, sticky="ew")
        outer.rowconfigure(0, weight=1)
        outer.columnconfigure(0, weight=1)

        self.overall_calendar_frame = tk.Frame(canvas, bg="white")
        canvas.create_window((0, 0), window=self.overall_calendar_frame, anchor="nw")
        self.overall_calendar_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )
        self.overall_calendar_canvas = canvas

    def buka_kalender_keseluruhan(self):
        """Buka kalender keseluruhan sebagai jendela global, bukan tab kegiatan."""
        if (
            getattr(self, "overall_calendar_window", None) is not None
            and self.overall_calendar_window.winfo_exists()
        ):
            self.overall_calendar_window.lift()
            self.overall_calendar_window.focus_force()
            self.load_kalender_keseluruhan()
            return

        win = tk.Toplevel(self)
        self.overall_calendar_window = win
        win.title("Kalender Keseluruhan")
        win.geometry("1500x850")
        win.minsize(1200, 650)
        win.transient(self)
        win.protocol("WM_DELETE_WINDOW", lambda: self.tutup_kalender_keseluruhan(win))

        self.build_kalender_keseluruhan(win)
        self.load_kalender_keseluruhan()

    def tutup_kalender_keseluruhan(self, win):
        try:
            if win.winfo_exists():
                win.destroy()
        finally:
            self.overall_calendar_window = None
            for attr in (
                "overall_calendar_frame",
                "overall_calendar_canvas",
                "overall_info_kalender",
                "overall_info_libur",
                "overall_legend_kegiatan",
                "overall_calendar_detail",
            ):
                if hasattr(self, attr):
                    try:
                        delattr(self, attr)
                    except Exception:
                        pass

    def buka_data_pegawai(self):
        """Buka data pegawai sebagai satu jendela global."""
        if getattr(self, "data_pegawai_window", None) is not None and self.data_pegawai_window.winfo_exists():
            self.data_pegawai_window.lift()
            self.data_pegawai_window.focus_force()
            self.refresh_data_pegawai()
            return

        win = tk.Toplevel(self)
        self.data_pegawai_window = win
        win.title("Data Pegawai")
        win.geometry("1450x750")
        win.minsize(1100, 600)
        win.transient(self)
        win.protocol("WM_DELETE_WINDOW", lambda: self.tutup_data_pegawai(win))
        self.build_pegawai(win)
        self.refresh_data_pegawai()

    def tutup_data_pegawai(self, win):
        try:
            if win.winfo_exists():
                win.destroy()
        finally:
            self.data_pegawai_window = None
            self.table_pegawai = None
            self.search_data_pegawai_var = None
            self.label_jumlah_pegawai = None
            self.tampilkan_nonaktif_var = None

    def build_pegawai(self, parent):

        tk.Label(
            parent,
            text="DATA PEGAWAI",
            bg="#F4F6F8",
            font=("Segoe UI", 18, "bold")
        ).pack(
            anchor="w",
            padx=20,
            pady=(15, 5)
        )

        self.label_jumlah_pegawai = tk.Label(
            parent,
            text="",
            bg="#F4F6F8",
            fg="#666666",
            font=("Segoe UI", 10)
        )
        self.label_jumlah_pegawai.pack(
            anchor="w",
            padx=20
        )

        # =====================================================
        # TOOLBAR DATA PEGAWAI
        # =====================================================
        toolbar = tk.Frame(parent, bg="#F4F6F8")
        toolbar.pack(fill="x", padx=20, pady=(12, 5))

        ttk.Button(
            toolbar,
            text="＋ Tambah Pegawai",
            command=lambda: self.buka_form_pegawai()
        ).pack(side="left", padx=(0, 6))

        ttk.Button(
            toolbar,
            text="✎ Edit",
            command=self.edit_pegawai_terpilih
        ).pack(side="left", padx=6)

        ttk.Button(
            toolbar,
            text="🗑 Nonaktifkan",
            command=self.hapus_pegawai_terpilih
        ).pack(side="left", padx=6)

        ttk.Button(
            toolbar,
            text="🗑 Hapus Permanen",
            command=self.hapus_pegawai_permanen_terpilih
        ).pack(side="left", padx=6)

        ttk.Button(
            toolbar,
            text="↥ Impor dari Excel",
            command=self.impor_pegawai_dari_excel
        ).pack(side="left", padx=6)

        ttk.Button(
            toolbar,
            text="✓ Aktifkan",
            command=self.aktifkan_pegawai_terpilih
        ).pack(side="left", padx=6)

        ttk.Button(
            toolbar,
            text="↻ Refresh Semua",
            command=self.refresh_all
        ).pack(side="left", padx=6)

        self.tampilkan_nonaktif_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(
            toolbar,
            text="Tampilkan pegawai nonaktif",
            variable=self.tampilkan_nonaktif_var,
            command=self.refresh_data_pegawai
        ).pack(side="right", padx=(10, 0))

        # =====================================================
        # SEARCH DATA PEGAWAI
        # =====================================================
        search_frame = tk.Frame(parent, bg="#F4F6F8")
        search_frame.pack(fill="x", padx=20, pady=(5, 5))

        tk.Label(
            search_frame,
            text="Cari Pegawai:",
            bg="#F4F6F8",
            font=("Segoe UI", 10, "bold")
        ).pack(side="left", padx=(0, 8))

        self.search_data_pegawai_var = tk.StringVar()
        self.entry_search_data_pegawai = ttk.Entry(
            search_frame,
            textvariable=self.search_data_pegawai_var,
            width=45
        )
        self.entry_search_data_pegawai.pack(side="left")
        self.entry_search_data_pegawai.bind(
            "<KeyRelease>",
            lambda event: self.refresh_data_pegawai()
        )

        ttk.Button(
            search_frame,
            text="✕",
            width=3,
            command=self.clear_search_data_pegawai
        ).pack(side="left", padx=5)

        # =====================================================
        # TABLE
        # =====================================================
        frame = tk.Frame(parent, bg="white")
        frame.pack(fill="both", expand=True, padx=20, pady=15)

        columns = (
            "nama",
            "nip",
            "pangkat",
            "golongan",
            "jabatan",
            "tarif",
            "bulan_ini",
            "status"
        )

        self.table_pegawai = ttk.Treeview(
            frame,
            columns=columns,
            show="headings",
            selectmode="browse"
        )

        heads = {
            "nama": "Nama",
            "nip": "NIP",
            "pangkat": "Pangkat",
            "golongan": "Golongan",
            "jabatan": "Jabatan",
            "tarif": "Tarif",
            "bulan_ini": "Turun Bulan Ini",
            "status": "Status"
        }

        widths = {
            "nama": 280,
            "nip": 185,
            "pangkat": 170,
            "golongan": 105,
            "jabatan": 230,
            "tarif": 135,
            "bulan_ini": 145,
            "status": 100
        }

        for col in columns:
            self.table_pegawai.heading(col, text=heads[col])
            self.table_pegawai.column(
                col,
                width=widths[col],
                anchor="center"
            )

        self.table_pegawai.column("nama", anchor="w")
        self.table_pegawai.column("jabatan", anchor="w")

        self.table_pegawai.pack(side="left", fill="both", expand=True)

        scroll = ttk.Scrollbar(
            frame,
            orient="vertical",
            command=self.table_pegawai.yview
        )
        scroll.pack(side="right", fill="y")
        self.table_pegawai.configure(yscrollcommand=scroll.set)

        # Double-click baris = edit.
        self.table_pegawai.bind(
            "<Double-1>",
            lambda event: self.edit_pegawai_terpilih()
        )
        self.table_pegawai.bind(
            "<Return>",
            lambda event: self.edit_pegawai_terpilih()
        )
        self.table_pegawai.bind(
            "<Delete>",
            lambda event: self.hapus_pegawai_terpilih()
        )

        self.refresh_data_pegawai()

    def _pegawai_terpilih_id(self):
        """Mengambil ID pegawai dari baris yang sedang dipilih."""
        if not hasattr(self, "table_pegawai") or self.table_pegawai is None:
            return None
        try:
            pilihan = self.table_pegawai.selection()
        except tk.TclError:
            return None
        if not pilihan:
            return None
        try:
            return int(pilihan[0])
        except (ValueError, TypeError):
            return None

    def _ambil_pegawai_by_id(self, pegawai_id):
        conn = get_connection()
        try:
            return conn.execute(
                """
                SELECT id, nama, nip, pangkat, golongan_asli,
                       golongan_utama, jabatan, tarif, aktif
                FROM pegawai
                WHERE id = ?
                """,
                (pegawai_id,)
            ).fetchone()
        finally:
            conn.close()

    def edit_pegawai_terpilih(self):
        pegawai_id = self._pegawai_terpilih_id()
        if pegawai_id is None:
            messagebox.showinfo(
                "Data Pegawai",
                "Pilih satu pegawai yang ingin diedit terlebih dahulu.",
                parent=getattr(self, "data_pegawai_window", self)
            )
            return
        self.buka_form_pegawai(pegawai_id)

    def buka_form_pegawai(self, pegawai_id=None):
        """Form tambah/edit pegawai."""
        parent = getattr(self, "data_pegawai_window", self)
        edit_mode = pegawai_id is not None

        data = None
        if edit_mode:
            data = self._ambil_pegawai_by_id(pegawai_id)
            if not data:
                messagebox.showerror(
                    "Data Pegawai",
                    "Data pegawai tidak ditemukan.",
                    parent=parent
                )
                self.refresh_data_pegawai()
                return

        win = tk.Toplevel(parent)
        win.title("Edit Pegawai" if edit_mode else "Tambah Pegawai")
        win.geometry("560x560")
        win.minsize(500, 500)
        win.transient(parent)
        win.grab_set()

        outer = tk.Frame(win, bg="#F4F6F8")
        outer.pack(fill="both", expand=True)

        tk.Label(
            outer,
            text="EDIT DATA PEGAWAI" if edit_mode else "TAMBAH DATA PEGAWAI",
            bg="#F4F6F8",
            font=("Segoe UI", 16, "bold")
        ).pack(anchor="w", padx=25, pady=(20, 5))

        tk.Label(
            outer,
            text="Isi data pegawai. Tarif disimpan sebagai tarif aktual pegawai.",
            bg="#F4F6F8",
            fg="#666666",
            font=("Segoe UI", 9)
        ).pack(anchor="w", padx=25, pady=(0, 15))

        form = tk.Frame(outer, bg="white", bd=1, relief="solid")
        form.pack(fill="both", expand=True, padx=25, pady=(0, 15))

        values = {
            "nama": data[1] if data else "",
            "nip": data[2] if data else "",
            "pangkat": data[3] if data else "",
            "golongan": data[4] if data else "",
            "jabatan": data[6] if data else "",
            "tarif": str(data[7]) if data else "",
        }

        vars_ = {key: tk.StringVar(value=value or "") for key, value in values.items()}

        fields = [
            ("Nama *", "nama"),
            ("NIP", "nip"),
            ("Pangkat", "pangkat"),
            ("Golongan", "golongan"),
            ("Jabatan", "jabatan"),
            ("Tarif *", "tarif"),
        ]

        entries = {}
        for row, (label, key) in enumerate(fields):
            tk.Label(
                form,
                text=label,
                bg="white",
                font=("Segoe UI", 10, "bold")
            ).grid(row=row, column=0, sticky="w", padx=18, pady=9)

            ent = ttk.Entry(form, textvariable=vars_[key], width=45)
            ent.grid(row=row, column=1, sticky="ew", padx=(0, 18), pady=9)
            entries[key] = ent

        form.columnconfigure(1, weight=1)

        # Golongan utama hanya informasi turunan; pengguna tetap mengisi
        # golongan asli seperti III d, III c, VII, X, dst.
        tk.Label(
            form,
            text="Golongan Utama",
            bg="white",
            fg="#666666",
            font=("Segoe UI", 9)
        ).grid(row=len(fields), column=0, sticky="w", padx=18, pady=(5, 14))

        gol_utama_var = tk.StringVar(value=golongan_utama_dari(vars_["golongan"].get()))
        tk.Label(
            form,
            textvariable=gol_utama_var,
            bg="white",
            fg="#333333",
            font=("Segoe UI", 10, "bold")
        ).grid(row=len(fields), column=1, sticky="w", padx=(0, 18), pady=(5, 14))

        def update_golongan_utama(*_):
            gol_utama_var.set(golongan_utama_dari(vars_["golongan"].get()))

        vars_["golongan"].trace_add("write", update_golongan_utama)

        def simpan():
            nama = vars_["nama"].get().strip()
            nip = vars_["nip"].get().strip()
            pangkat = vars_["pangkat"].get().strip()
            golongan = vars_["golongan"].get().strip()
            jabatan = vars_["jabatan"].get().strip()
            tarif_text = vars_["tarif"].get().strip()

            if not nama:
                messagebox.showwarning(
                    "Data Pegawai",
                    "Nama pegawai wajib diisi.",
                    parent=win
                )
                entries["nama"].focus_set()
                return

            try:
                tarif = parse_rupiah(tarif_text)
            except (ValueError, TypeError):
                messagebox.showwarning(
                    "Data Pegawai",
                    "Tarif harus berupa angka, misalnya 80000 atau Rp 80.000.",
                    parent=win
                )
                entries["tarif"].focus_set()
                return

            if tarif <= 0:
                messagebox.showwarning(
                    "Data Pegawai",
                    "Tarif harus lebih besar dari 0.",
                    parent=win
                )
                entries["tarif"].focus_set()
                return

            # NIP tidak boleh dobel jika diisi.
            conn = get_connection()
            try:
                if nip:
                    if edit_mode:
                        duplikat = conn.execute(
                            "SELECT id FROM pegawai WHERE nip = ? AND id <> ? LIMIT 1",
                            (nip, pegawai_id)
                        ).fetchone()
                    else:
                        duplikat = conn.execute(
                            "SELECT id FROM pegawai WHERE nip = ? LIMIT 1",
                            (nip,)
                        ).fetchone()
                    if duplikat:
                        messagebox.showwarning(
                            "Data Pegawai",
                            f"NIP {nip} sudah digunakan oleh pegawai lain.",
                            parent=win
                        )
                        entries["nip"].focus_set()
                        return

                golongan_utama = golongan_utama_dari(golongan)

                if edit_mode:
                    conn.execute(
                        """
                        UPDATE pegawai
                        SET nama = ?, nip = ?, pangkat = ?,
                            golongan_asli = ?, golongan_utama = ?,
                            jabatan = ?, tarif = ?
                        WHERE id = ?
                        """,
                        (
                            nama, nip, pangkat,
                            golongan, golongan_utama,
                            jabatan, tarif, pegawai_id
                        )
                    )
                else:
                    no_row = conn.execute(
                        "SELECT COALESCE(MAX(no), 0) + 1 FROM pegawai"
                    ).fetchone()
                    no = int(no_row[0] or 1)
                    conn.execute(
                        """
                        INSERT INTO pegawai
                        (no, nama, nip, pangkat, golongan_asli,
                         golongan_utama, jabatan, tarif, aktif)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, 1)
                        """,
                        (
                            no, nama, nip, pangkat,
                            golongan, golongan_utama,
                            jabatan, tarif
                        )
                    )

                conn.commit()
            except Exception as exc:
                conn.rollback()
                messagebox.showerror(
                    "Data Pegawai",
                    f"Gagal menyimpan data pegawai.\n\n{exc}",
                    parent=win
                )
                return
            finally:
                conn.close()

            win.grab_release()
            win.destroy()
            self.pegawai = load_pegawai()
            self.refresh_pj_options()
            self.refresh_selection_table()
            self.load_bagian_data()
            self.refresh_data_pegawai()
            self.load_kalender()
            self.load_kalender_keseluruhan()

        button_frame = tk.Frame(outer, bg="#F4F6F8")
        button_frame.pack(fill="x", padx=25, pady=(0, 20))

        ttk.Button(
            button_frame,
            text="Batal",
            command=lambda: (win.grab_release(), win.destroy())
        ).pack(side="right", padx=(8, 0))
        ttk.Button(
            button_frame,
            text="Simpan",
            command=simpan
        ).pack(side="right")

        entries["nama"].focus_set()
        win.bind("<Escape>", lambda event: (win.grab_release(), win.destroy()))
        win.bind("<Return>", lambda event: simpan())

    def aktifkan_pegawai_terpilih(self):
        pegawai_id = self._pegawai_terpilih_id()
        if pegawai_id is None:
            messagebox.showinfo(
                "Data Pegawai",
                "Pilih satu pegawai yang ingin diaktifkan kembali.",
                parent=getattr(self, "data_pegawai_window", self)
            )
            return

        data = self._ambil_pegawai_by_id(pegawai_id)
        if not data:
            return
        if int(data[8] or 0) == 1:
            messagebox.showinfo(
                "Data Pegawai",
                f"{data[1]} sudah berstatus aktif.",
                parent=getattr(self, "data_pegawai_window", self)
            )
            return

        jawab = messagebox.askyesno(
            "Aktifkan Pegawai",
            f"Aktifkan kembali pegawai {data[1]}?",
            parent=getattr(self, "data_pegawai_window", self)
        )
        if not jawab:
            return

        conn = get_connection()
        try:
            conn.execute("UPDATE pegawai SET aktif = 1 WHERE id = ?", (pegawai_id,))
            conn.commit()
        except Exception as exc:
            conn.rollback()
            messagebox.showerror(
                "Aktifkan Pegawai",
                f"Gagal mengaktifkan pegawai.\n\n{exc}",
                parent=getattr(self, "data_pegawai_window", self)
            )
            return
        finally:
            conn.close()

        self.pegawai = load_pegawai()
        self.refresh_pj_options()
        self.refresh_selection_table()
        self.load_bagian_data()
        self.refresh_data_pegawai()
        self.load_kalender()
        self.load_kalender_keseluruhan()

    def hapus_pegawai_terpilih(self):
        pegawai_id = self._pegawai_terpilih_id()
        if pegawai_id is None:
            messagebox.showinfo(
                "Data Pegawai",
                "Pilih satu pegawai yang ingin dihapus terlebih dahulu.",
                parent=getattr(self, "data_pegawai_window", self)
            )
            return

        data = self._ambil_pegawai_by_id(pegawai_id)
        if not data:
            return

        nama = data[1] or "Tanpa Nama"

        # Jangan menonaktifkan pegawai yang masih dipakai pada kegiatan
        # periode aktif. Pengguna perlu melepas/mengganti pegawai tersebut
        # dari kegiatan bulan berjalan terlebih dahulu agar jadwal tetap valid.
        conn = get_connection()
        try:
            dipakai_periode_ini = conn.execute(
                """
                SELECT COUNT(*)
                FROM anggota_bagian ab
                JOIN bagian b ON b.id = ab.bagian_id
                WHERE b.periode_id = ? AND ab.pegawai_id = ?
                """,
                (self.periode_id, pegawai_id)
            ).fetchone()[0]
        finally:
            conn.close()

        if dipakai_periode_ini:
            messagebox.showwarning(
                "Pegawai Masih Dipakai",
                f"{nama} masih terdaftar pada kegiatan di {NAMA_BULAN[self.bulan]} {self.tahun}.\n\n"
                "Lepaskan pegawai dari kegiatan bulan tersebut terlebih dahulu, "
                "baru pegawai dapat dinonaktifkan.",
                parent=getattr(self, "data_pegawai_window", self)
            )
            return

        # Hapus secara soft-delete supaya riwayat penugasan lama tetap aman.
        jawab = messagebox.askyesno(
            "Hapus Pegawai",
            f"Yakin ingin menghapus/nonaktifkan pegawai berikut?\n\n"
            f"Nama : {nama}\n"
            f"NIP  : {data[2] or '-'}\n\n"
            f"Data tidak dihapus permanen. Pegawai akan dibuat nonaktif "
            f"agar riwayat penugasan tetap tersimpan.",
            parent=getattr(self, "data_pegawai_window", self)
        )
        if not jawab:
            return

        conn = get_connection()
        try:
            conn.execute(
                "UPDATE pegawai SET aktif = 0 WHERE id = ?",
                (pegawai_id,)
            )
            conn.commit()
        except Exception as exc:
            conn.rollback()
            messagebox.showerror(
                "Hapus Pegawai",
                f"Gagal menonaktifkan pegawai.\n\n{exc}",
                parent=getattr(self, "data_pegawai_window", self)
            )
            return
        finally:
            conn.close()

        self.pegawai = load_pegawai()
        self.refresh_pj_options()
        self.refresh_selection_table()
        self.load_bagian_data()
        self.refresh_data_pegawai()
        self.load_kalender()
        self.load_kalender_keseluruhan()


    def hapus_pegawai_permanen_terpilih(self):
        """Menghapus pegawai secara permanen beserta seluruh riwayat yang terkait."""
        pegawai_id = self._pegawai_terpilih_id()
        parent = getattr(self, "data_pegawai_window", self)
        if pegawai_id is None:
            messagebox.showinfo(
                "Hapus Permanen",
                "Pilih satu pegawai yang ingin dihapus permanen terlebih dahulu.",
                parent=parent
            )
            return

        data = self._ambil_pegawai_by_id(pegawai_id)
        if not data:
            return

        nama = data[1] or "Tanpa Nama"
        nip = data[2] or "-"

        conn = get_connection()
        try:
            jumlah_jadwal = conn.execute(
                "SELECT COUNT(*) FROM jadwal WHERE pegawai_id = ?",
                (pegawai_id,)
            ).fetchone()[0]
            jumlah_anggota = conn.execute(
                "SELECT COUNT(*) FROM anggota_bagian WHERE pegawai_id = ?",
                (pegawai_id,)
            ).fetchone()[0]
        finally:
            conn.close()

        pesan = (
            f"PERINGATAN: tindakan ini PERMANEN.\n\n"
            f"Pegawai:\n{nama}\nNIP: {nip}\n\n"
            f"Yang akan dihapus:\n"
            f"• Data master pegawai\n"
            f"• {jumlah_anggota} keanggotaan kegiatan\n"
            f"• {jumlah_jadwal} jadwal/rekam turun\n\n"
            f"Riwayat pegawai ini tidak dapat dipulihkan dari aplikasi.\n"
            f"Pastikan sudah melakukan backup/export terlebih dahulu.\n\n"
            f"Lanjutkan?"
        )
        if not messagebox.askyesno("Hapus Pegawai Permanen", pesan, parent=parent):
            return

        konfirmasi = simpledialog.askstring(
            "Konfirmasi Terakhir",
            f"Untuk benar-benar menghapus {nama}, ketik:\n\nHAPUS PERMANEN",
            parent=parent
        )
        if konfirmasi != "HAPUS PERMANEN":
            messagebox.showinfo(
                "Dibatalkan",
                "Penghapusan permanen dibatalkan. Data pegawai tetap aman.",
                parent=parent
            )
            return

        conn = get_connection()
        try:
            # Lepaskan dulu referensi PJ agar aman jika foreign key aktif.
            conn.execute(
                "UPDATE bagian SET penanggung_jawab_id = NULL WHERE penanggung_jawab_id = ?",
                (pegawai_id,)
            )
            conn.execute("DELETE FROM jadwal WHERE pegawai_id = ?", (pegawai_id,))
            conn.execute("DELETE FROM anggota_bagian WHERE pegawai_id = ?", (pegawai_id,))
            conn.execute("DELETE FROM pegawai WHERE id = ?", (pegawai_id,))
            conn.commit()
        except Exception as exc:
            conn.rollback()
            messagebox.showerror(
                "Hapus Permanen Gagal",
                f"Pegawai gagal dihapus permanen.\n\n{exc}",
                parent=parent
            )
            return
        finally:
            conn.close()

        self.pegawai = load_pegawai()
        self.selected_ids.discard(pegawai_id)
        self.refresh_pj_options()
        self.refresh_selection_table()
        self.load_bagian_data()
        self.refresh_data_pegawai()
        self.load_kalender()
        self.load_kalender_keseluruhan()
        self.refresh_data_kegiatan()
        messagebox.showinfo(
            "Berhasil",
            f"Pegawai {nama} sudah dihapus permanen beserta seluruh riwayat terkait.",
            parent=parent
        )

    def impor_pegawai_dari_excel(self):
        """
        Mengganti seluruh master pegawai dari file Excel yang dipilih user.

        Nama file bebas. Syarat satu-satunya: workbook memiliki sheet
        bernama persis ``DaftarPegawai``. Kolom dicari berdasarkan nama
        header, bukan posisi kolom, sehingga sheet tetap bisa memiliki
        kolom tambahan.
        """
        parent = getattr(self, "data_pegawai_window", self)
        if load_workbook is None:
            messagebox.showerror(
                "Import Excel",
                "Library openpyxl belum terpasang. Jalankan: pip install openpyxl",
                parent=parent
            )
            return

        path = filedialog.askopenfilename(
            parent=parent,
            title="Pilih File Excel Pegawai",
            filetypes=[
                ("Excel Macro-Enabled (*.xlsm)", "*.xlsm"),
                ("Excel Workbook (*.xlsx)", "*.xlsx"),
                ("Semua File Excel", "*.xlsm *.xlsx"),
            ]
        )
        if not path:
            return

        # Nama file TIDAK diperiksa. Yang diperiksa adalah nama sheet.
        try:
            wb = load_workbook(
                filename=path,
                read_only=True,
                data_only=True,
                keep_vba=path.lower().endswith(".xlsm")
            )
        except Exception as exc:
            messagebox.showerror(
                "Import Excel",
                "File Excel tidak dapat dibaca.\n\n"
                f"{exc}",
                parent=parent
            )
            return

        sheet_target = "DaftarPegawai"
        if sheet_target not in wb.sheetnames:
            tersedia = "\n".join(f"• {nama}" for nama in wb.sheetnames[:30])
            if len(wb.sheetnames) > 30:
                tersedia += "\n• ..."
            wb.close()
            messagebox.showerror(
                "Import Excel",
                f"Sheet '{sheet_target}' tidak ditemukan.\n\n"
                "Nama file bebas, tetapi workbook wajib memiliki sheet "
                "bernama persis 'DaftarPegawai'.\n\n"
                "Sheet yang ditemukan:\n"
                f"{tersedia}",
                parent=parent
            )
            return

        ws = wb[sheet_target]

        # ---------------------------------------------------------
        # Cari header berdasarkan NAMA KOLOM, bukan posisi kolom.
        # Contoh file kamu:
        # NO | NAMA | NIP | PANGKAT | GOL | JABATAN | Angka | ...
        # ---------------------------------------------------------
        def norm_header(value):
            if value is None:
                return ""
            return re.sub(r"\s+", " ", str(value).strip().upper())

        header_row_no = None
        header_map = {}
        nama_alias = {"NAMA", "NAMA PEGAWAI"}
        nip_alias = {"NIP"}
        pangkat_alias = {"PANGKAT"}
        gol_alias = {"GOL", "GOLONGAN", "GOLONGAN ASLI"}
        jabatan_alias = {"JABATAN"}
        tarif_alias = {"ANGKA", "TARIF", "NOMINAL", "ANGKA SESUAI INPUT"}

        max_scan_rows = min(ws.max_row or 1, 30)
        for r in range(1, max_scan_rows + 1):
            kandidat = {}
            for c in range(1, (ws.max_column or 1) + 1):
                h = norm_header(ws.cell(r, c).value)
                if h:
                    kandidat[h] = c

            if (
                any(h in kandidat for h in nama_alias)
                and any(h in kandidat for h in nip_alias)
                and any(h in kandidat for h in pangkat_alias)
                and any(h in kandidat for h in gol_alias)
                and any(h in kandidat for h in jabatan_alias)
            ):
                header_row_no = r
                header_map = kandidat
                break

        if header_row_no is None:
            wb.close()
            messagebox.showerror(
                "Import Excel",
                "Sheet 'DaftarPegawai' ditemukan, tetapi header pegawai tidak dikenali.\n\n"
                "Header minimal yang harus ada:\n"
                "NAMA, NIP, PANGKAT, GOL, JABATAN\n\n"
                "Tarif boleh menggunakan salah satu nama: Angka, Tarif, Nominal, "
                "atau Angka Sesuai Input.",
                parent=parent
            )
            return

        def find_col(aliases):
            for alias in aliases:
                if alias in header_map:
                    return header_map[alias]
            return None

        col_nama = find_col(nama_alias)
        col_nip = find_col(nip_alias)
        col_pangkat = find_col(pangkat_alias)
        col_gol = find_col(gol_alias)
        col_jabatan = find_col(jabatan_alias)
        # Prioritaskan "ANGKA" seperti file DaftarPegawai kamu.
        col_tarif = (
            header_map.get("ANGKA")
            or header_map.get("TARIF")
            or header_map.get("NOMINAL")
            or header_map.get("ANGKA SESUAI INPUT")
        )
        col_no = header_map.get("NO")

        rows = []
        nip_seen = set()
        errors = []

        for row_no in range(header_row_no + 1, (ws.max_row or 0) + 1):
            def cell_value(col):
                return ws.cell(row_no, col).value if col else None

            values = [cell_value(c) for c in [col_no, col_nama, col_nip, col_pangkat, col_gol, col_jabatan, col_tarif]]
            if not any(v is not None and str(v).strip() != "" for v in values):
                continue

            no = nilai_excel_teks(cell_value(col_no))
            nama = nilai_excel_teks(cell_value(col_nama))
            nip = nilai_excel_teks(cell_value(col_nip))
            pangkat = nilai_excel_teks(cell_value(col_pangkat))
            golongan = nilai_excel_teks(cell_value(col_gol))
            jabatan = nilai_excel_teks(cell_value(col_jabatan))
            tarif = parse_tarif_excel(cell_value(col_tarif))

            # Jika baris tidak punya nama sama sekali, abaikan baris kosong/teks tambahan.
            if not nama:
                continue

            if nip and nip in nip_seen:
                errors.append(f"Baris {row_no}: NIP {nip} duplikat")
                continue
            if nip:
                nip_seen.add(nip)

            rows.append({
                "no": no,
                "nama": nama,
                "nip": nip,
                "pangkat": pangkat,
                "golongan": golongan,
                "golongan_utama": golongan_utama_dari(golongan),
                "jabatan": jabatan,
                "tarif": tarif,
            })

        wb.close()

        if errors:
            messagebox.showerror(
                "Import Excel Dibatalkan",
                "Ada masalah pada data Excel. Tidak ada data yang diubah.\n\n"
                + "\n".join(errors[:20])
                + ("\n..." if len(errors) > 20 else ""),
                parent=parent
            )
            return

        if not rows:
            messagebox.showwarning(
                "Import Excel",
                "Sheet 'DaftarPegawai' ditemukan, tetapi tidak ada data pegawai yang terbaca.",
                parent=parent
            )
            return

        tanpa_tarif = sum(1 for r in rows if r["tarif"] <= 0)
        peringatan_tarif = (
            f"\n\nPERHATIAN: {tanpa_tarif} pegawai tidak memiliki tarif valid. "
            "Tarif master berdasarkan golongan akan digunakan jika tersedia."
            if tanpa_tarif else ""
        )

        jawab = messagebox.askyesno(
            "Ganti Seluruh Data Pegawai",
            f"File: {path}\n"
            f"Sheet: {sheet_target}\n"
            f"Header ditemukan di baris: {header_row_no}\n"
            f"Pegawai terbaca: {len(rows)} orang\n\n"
            "IMPORT INI AKAN MENGGANTIKAN SELURUH DATA PEGAWAI LAMA.\n"
            "Semua keanggotaan dan jadwal yang terkait pegawai lama juga akan dihapus.\n"
            "Data kegiatan/periode tetap ada, tetapi anggotanya akan kosong.\n\n"
            "Pastikan data lama sudah di-backup/export sebelum melanjutkan."
            + peringatan_tarif
            + "\n\nLanjutkan?",
            parent=parent
        )
        if not jawab:
            return

        konfirmasi = simpledialog.askstring(
            "Konfirmasi Penggantian Data",
            "Ketik GANTI untuk menghapus data pegawai lama dan memasukkan data Excel baru:",
            parent=parent
        )
        if konfirmasi != "GANTI":
            messagebox.showinfo(
                "Dibatalkan",
                "Import dibatalkan. Tidak ada data yang diubah.",
                parent=parent
            )
            return

        conn = get_connection()
        try:
            # Hapus relasi lama terlebih dahulu. Master pegawai baru kemudian dimasukkan.
            conn.execute("UPDATE bagian SET penanggung_jawab_id = NULL")
            conn.execute("DELETE FROM jadwal")
            conn.execute("DELETE FROM anggota_bagian")
            conn.execute("DELETE FROM pegawai")

            tarif_master = {
                str(r[0]).strip().upper(): int(r[1] or 0)
                for r in conn.execute("SELECT golongan, nominal FROM tarif").fetchall()
            }

            for idx, item in enumerate(rows, start=1):
                tarif = item["tarif"]
                if tarif <= 0:
                    tarif = tarif_master.get(item["golongan_utama"], 0)

                conn.execute(
                    """
                    INSERT INTO pegawai
                    (no, nama, nip, pangkat, golongan_asli, golongan_utama,
                     jabatan, tarif, aktif)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, 1)
                    """,
                    (
                        idx,
                        item["nama"],
                        item["nip"],
                        item["pangkat"],
                        item["golongan"],
                        item["golongan_utama"],
                        item["jabatan"],
                        tarif
                    )
                )
            conn.commit()
        except Exception as exc:
            conn.rollback()
            messagebox.showerror(
                "Import Excel Gagal",
                f"Import gagal dan perubahan dibatalkan.\n\n{exc}",
                parent=parent
            )
            return
        finally:
            conn.close()

        # Semua tampilan di-refresh dari database setelah import.
        # Jangan memanggil refresh satu per satu di sini karena itu membuat
        # kalender/tabel dibangun berulang kali dan terasa berat.
        self.selected_ids.clear()
        self.refresh_all()

        messagebox.showinfo(
            "Import Berhasil",
            f"Berhasil mengimpor {len(rows)} pegawai dari sheet '{sheet_target}'.\n\n"
            "Nama file bebas; yang diwajibkan adalah nama sheet 'DaftarPegawai'.\n"
            "Keanggotaan dan jadwal lama sudah dikosongkan. Data kegiatan/periode tetap ada.",
            parent=parent
        )

    def refresh_data_pegawai(self):

        try:
            self.pegawai = load_pegawai()
        except Exception:
            pass

        if not hasattr(self, "table_pegawai") or self.table_pegawai is None:
            return

        try:
            if not self.table_pegawai.winfo_exists():
                self.table_pegawai = None
                return
        except tk.TclError:
            self.table_pegawai = None
            return

        try:
            for item in self.table_pegawai.get_children():
                self.table_pegawai.delete(item)
        except tk.TclError:
            self.table_pegawai = None
            return

        if hasattr(self, "label_jumlah_pegawai"):
            try:
                self.label_jumlah_pegawai.config(
                    text=f"{len(self.pegawai)} pegawai aktif"
                )
            except tk.TclError:
                pass

        riwayat = load_riwayat_turun(self.periode_id)

        keyword = ""
        if hasattr(self, "search_data_pegawai_var") and self.search_data_pegawai_var:
            keyword = self.search_data_pegawai_var.get().strip().lower()

        tampil_nonaktif = bool(
            getattr(self, "tampilkan_nonaktif_var", tk.BooleanVar(value=False)).get()
        )

        self.table_pegawai.tag_configure("baris1", background="#FFFFFF")
        self.table_pegawai.tag_configure("baris2", background="#F1F4F7")
        self.table_pegawai.tag_configure("nonaktif", foreground="#888888", background="#F4F4F4")

        # Ambil semua pegawai agar checkbox "Tampilkan nonaktif" benar-benar
        # dapat menampilkan pegawai yang sudah dinonaktifkan.
        conn = get_connection()
        try:
            semua_pegawai = conn.execute(
                """
                SELECT id, nama, nip, pangkat, golongan_asli,
                       golongan_utama, jabatan, tarif, aktif
                FROM pegawai
                ORDER BY nama
                """
            ).fetchall()
        finally:
            conn.close()

        nomor_baris = 0
        for p in semua_pegawai:
            aktif = int(p[8] or 0) == 1
            if not aktif and not tampil_nonaktif:
                continue

            teks_pencarian = " ".join([
                str(p[1] or ""),
                str(p[2] or ""),
                str(p[3] or ""),
                str(p[4] or ""),
                str(p[5] or ""),
                str(p[6] or "")
            ]).lower()

            if keyword and keyword not in teks_pencarian:
                continue

            pegawai_id = p[0]
            data = riwayat.get(
                pegawai_id,
                {"bulan_ini": 0, "total": 0}
            )

            tag = "baris1" if nomor_baris % 2 == 0 else "baris2"
            nomor_baris += 1
            if not aktif:
                tag = "nonaktif"

            self.table_pegawai.insert(
                "",
                "end",
                iid=str(pegawai_id),
                values=(
                    p[1],
                    p[2],
                    p[3],
                    p[4],
                    p[6],
                    rupiah(p[7]),
                    f"{data['bulan_ini']} kali",
                    "Aktif" if aktif else "Nonaktif"
                ),
                tags=(tag,)
            )

    # =====================================================
    # EXPORT EXCEL
    # =====================================================

    @staticmethod
    def _excel_sheet_name(name, fallback="Sheet"):
        """Membersihkan nama sheet Excel (maks. 31 karakter)."""
        name = str(name or "").strip()
        name = re.sub(r'[\\/*?:\[\]]', '-', name)
        name = name or fallback
        return name[:31]

    @staticmethod
    def _excel_unique_sheet_name(wb, name):
        """Memastikan nama worksheet unik dan tetap <= 31 karakter."""
        base = App._excel_sheet_name(name)
        if base not in wb.sheetnames:
            return base
        nomor = 2
        while True:
            suffix = f" ({nomor})"
            kandidat = base[:31 - len(suffix)] + suffix
            if kandidat not in wb.sheetnames:
                return kandidat
            nomor += 1

    @staticmethod
    def _format_excel_sheet(ws, header_row=None, widths=None):
        """Format umum workbook hasil export."""
        thin = Side(style="thin", color="D9D9D9")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        ws.sheet_view.showGridLines = False

        if header_row:
            for cell in ws[header_row]:
                if cell.value is not None:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="17324D")
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.border = border
            ws.row_dimensions[header_row].height = 28
            ws.freeze_panes = f"A{header_row + 1}"
            ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"

        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cell.border = border
                    if cell.row != header_row:
                        cell.alignment = Alignment(vertical="center", wrap_text=True)

        if widths:
            for col, width in widths.items():
                ws.column_dimensions[col].width = width

    def _export_calendar_sheet(self, wb, conn, bulan, tahun, periode_id):
        ws = wb.active
        ws.title = "Kalender Keseluruhan"
        ws.sheet_view.zoomScale = 80

        jumlah_hari = calendar.monthrange(tahun, bulan)[1]
        hari_libur = load_hari_libur(periode_id)

        kegiatan_rows = conn.execute(
            """
            SELECT id, kode, COALESCE(nama_bagian, ''), COALESCE(kegiatan, ''),
                   COALESCE(lokasi, ''), COALESCE(dana, 0), penanggung_jawab_id
            FROM bagian
            WHERE periode_id = ?
            ORDER BY id
            """,
            (periode_id,)
        ).fetchall()

        jadwal_rows = conn.execute(
            """
            SELECT j.pegawai_id, j.hari, b.kode, COALESCE(b.nama_bagian, '')
            FROM jadwal j
            JOIN bagian b ON b.id = j.bagian_id
            WHERE j.periode_id = ?
            ORDER BY j.pegawai_id, j.hari
            """,
            (periode_id,)
        ).fetchall()

        pegawai_rows = conn.execute(
            """
            SELECT id, nama, nip, pangkat, golongan_asli, jabatan, tarif
            FROM pegawai
            WHERE aktif = 1
            ORDER BY nama
            """
        ).fetchall()

        nama_kegiatan = {
            row[1]: (row[2].strip() if row[2] else "") or f"Kegiatan {row[1]}"
            for row in kegiatan_rows
        }
        jadwal_map = {(r[0], r[1]): (r[2], nama_kegiatan.get(r[2], f"Kegiatan {r[2]}")) for r in jadwal_rows}

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=jumlah_hari + 1)
        ws.cell(1, 1, f"PENUGASAN {NAMA_BULAN[bulan].upper()} {tahun}")
        ws.cell(1, 1).font = Font(size=16, bold=True, color="FFFFFF")
        ws.cell(1, 1).fill = PatternFill("solid", fgColor="17324D")
        ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 32

        ws.cell(2, 1, "Nama Pegawai")
        ws.cell(2, 1).font = Font(bold=True, color="FFFFFF")
        ws.cell(2, 1).fill = PatternFill("solid", fgColor="17324D")
        ws.cell(2, 1).alignment = Alignment(horizontal="center", vertical="center")

        nama_hari = ["Sen", "Sel", "Rab", "Kam", "Jum", "Sab", "Min"]
        for hari in range(1, jumlah_hari + 1):
            tanggal = datetime(tahun, bulan, hari)
            ws.cell(2, hari + 1, f"{hari}\n{nama_hari[tanggal.weekday()]}")
            ws.cell(2, hari + 1).font = Font(bold=True, color="FFFFFF")
            ws.cell(2, hari + 1).fill = PatternFill("solid", fgColor="17324D")
            ws.cell(2, hari + 1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[2].height = 32

        palette = ["B8D4EF", "C7E6C5", "F3D6A5", "D9C2E9", "F4B8B8", "BFE3E3", "E7D6A8", "C9D7B8", "D5C4A1", "C7D4F2", "E5C1CD", "CDE4C8"]
        fills = {row[1]: PatternFill("solid", fgColor=palette[i % len(palette)]) for i, row in enumerate(kegiatan_rows)}
        holiday_fill = PatternFill("solid", fgColor="ECECEC")
        thin = Side(style="thin", color="D9D9D9")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for r_idx, pegawai in enumerate(pegawai_rows, start=3):
            ws.cell(r_idx, 1, pegawai[1] or "Tanpa Nama")
            ws.cell(r_idx, 1).alignment = Alignment(vertical="center", wrap_text=True)
            ws.cell(r_idx, 1).border = border
            for hari in range(1, jumlah_hari + 1):
                cell = ws.cell(r_idx, hari + 1)
                tanggal = datetime(tahun, bulan, hari)
                detail = jadwal_map.get((pegawai[0], hari))
                if tanggal.weekday() == 6 or hari in hari_libur:
                    cell.value = "LIBUR"
                    cell.fill = holiday_fill
                elif detail:
                    cell.value = f"{detail[0]} - {detail[1]}"
                    cell.fill = fills.get(detail[0], PatternFill("solid", fgColor="FFFFFF"))
                else:
                    cell.value = ""
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border

        legend_start = len(pegawai_rows) + 5
        ws.cell(legend_start, 1, "Keterangan Kegiatan")
        ws.cell(legend_start, 1).font = Font(bold=True)
        for i, row in enumerate(kegiatan_rows, start=1):
            cell = ws.cell(legend_start + i, 1, f"{row[1]} - {nama_kegiatan[row[1]]}")
            cell.fill = fills[row[1]]
            cell.border = border
        ws.cell(legend_start + len(kegiatan_rows) + 1, 1, "LIBUR = Minggu atau hari libur manual")
        ws.cell(legend_start + len(kegiatan_rows) + 1, 1).fill = holiday_fill
        ws.freeze_panes = "B3"
        ws.column_dimensions["A"].width = 28
        for col in range(2, jumlah_hari + 2):
            ws.column_dimensions[get_column_letter(col)].width = 13
        ws.auto_filter.ref = f"A2:{get_column_letter(jumlah_hari + 1)}{max(2, len(pegawai_rows) + 2)}"
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.print_title_rows = "1:2"

    def _export_data_pegawai_sheet(self, wb, conn):
        ws = wb.create_sheet("Data Pegawai")
        ws.sheet_view.zoomScale = 90
        ws.append(["Nama", "NIP", "Pangkat", "Golongan", "Golongan Utama", "Jabatan", "Tarif", "Status", "Turun Bulan Ini"])

        riwayat = load_riwayat_turun(self.periode_id)
        rows = conn.execute(
            """
            SELECT id, nama, nip, pangkat, golongan_asli, golongan_utama, jabatan, tarif, aktif
            FROM pegawai
            ORDER BY nama
            """
        ).fetchall()
        for p in rows:
            r = riwayat.get(p[0], {"bulan_ini": 0})
            ws.append([
                p[1], p[2], p[3], p[4], p[5], p[6], p[7],
                "Aktif" if p[8] else "Nonaktif",
                r["bulan_ini"]
            ])
        for cell in ws[1]:
            cell.number_format = "@" if cell.column in (2,) else cell.number_format
        for row in ws.iter_rows(min_row=2, min_col=7, max_col=7):
            row[0].number_format = '#,##0'
        self._format_excel_sheet(ws, header_row=1, widths={"A": 30, "B": 20, "C": 18, "D": 12, "E": 16, "F": 28, "G": 14, "H": 12, "I": 18})

    def _export_kegiatan_sheet(self, wb, conn, bagian, bulan, tahun):
        bagian_id, kode, nama_bagian, deskripsi, lokasi, dana, pj_id = bagian
        judul = (nama_bagian or "").strip() or f"Kegiatan {kode}"
        sheet_name = self._excel_unique_sheet_name(wb, judul)
        ws = wb.create_sheet(sheet_name)
        ws.sheet_view.zoomScale = 90

        pj_row = conn.execute("SELECT nama FROM pegawai WHERE id = ?", (pj_id,)).fetchone() if pj_id else None
        pj_nama = pj_row[0] if pj_row else "-"

        metadata = [
            ("Kegiatan", judul),
            ("Kode", kode),
            ("Deskripsi", deskripsi or ""),
            ("Lokasi", lokasi or ""),
            ("Dana Bulanan", dana or 0),
            ("Penanggung Jawab", pj_nama),
            ("Periode", f"{NAMA_BULAN[bulan]} {tahun}"),
        ]
        for i, (label, value) in enumerate(metadata, start=1):
            ws.cell(i, 1, label).font = Font(bold=True)
            ws.cell(i, 2, value)
            if label == "Dana Bulanan":
                ws.cell(i, 2).number_format = '#,##0'
            ws.cell(i, 1).fill = PatternFill("solid", fgColor="D9E6F2")

        header_row = 9
        headers = ["No", "Nama", "NIP", "Pangkat", "Golongan", "Jabatan", "Tarif", "PJ", "Turun"]
        for c, h in enumerate(headers, start=1):
            ws.cell(header_row, c, h)

        ws.cell(header_row + 1, 1, "Catatan")
        ws.cell(header_row + 1, 2, "Kolom Turun hanya menghitung berapa kali pegawai turun pada kegiatan ini di bulan aktif. Jadwal lengkap dapat dilihat di sheet Kalender Keseluruhan.")
        ws.merge_cells(start_row=header_row + 1, start_column=2, end_row=header_row + 1, end_column=10)
        ws.cell(header_row + 1, 1).font = Font(bold=True)
        ws.cell(header_row + 1, 1).fill = PatternFill("solid", fgColor="FFF2CC")
        ws.cell(header_row + 1, 2).fill = PatternFill("solid", fgColor="FFF2CC")
        ws.cell(header_row + 1, 2).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[header_row + 1].height = 36

        anggota = conn.execute(
            """
            SELECT p.id, p.nama, p.nip, p.pangkat, p.golongan_asli, p.jabatan, p.tarif
            FROM anggota_bagian ab
            JOIN pegawai p ON p.id = ab.pegawai_id
            WHERE ab.bagian_id = ?
            ORDER BY p.nama
            """,
            (bagian_id,)
        ).fetchall()

        for nomor, p in enumerate(anggota, start=1):
            hari_rows = conn.execute(
                "SELECT hari FROM jadwal WHERE bagian_id = ? AND pegawai_id = ? ORDER BY hari",
                (bagian_id, p[0])
            ).fetchall()
            hari = [int(x[0]) for x in hari_rows]
            ws.append([
                nomor, p[1], p[2], p[3], p[4], p[5], p[6],
                "Ya" if p[0] == pj_id else "",
                len(hari)
            ])

        self._format_excel_sheet(
            ws,
            header_row=header_row,
            widths={"A": 7, "B": 30, "C": 20, "D": 18, "E": 12, "F": 28, "G": 14, "H": 8, "I": 12}
        )
        ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=10)
        ws.merge_cells(start_row=4, start_column=2, end_row=4, end_column=10)
        ws.row_dimensions[3].height = 45
        ws.row_dimensions[4].height = 30
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.print_title_rows = f"1:{header_row}"

    def export_penugasan_excel(self):
        """Export satu periode lengkap ke Excel: kalender, master pegawai, dan tiap kegiatan."""
        if Workbook is None:
            messagebox.showerror(
                "Export Excel",
                "Library openpyxl belum terpasang. Jalankan: pip install openpyxl",
                parent=self
            )
            return

        try:
            # NAMA_BULAN memakai indeks kalender asli (1=Januari, ..., 12=Desember)
            # sehingga tidak boleh ditambah 1 lagi.
            bulan = NAMA_BULAN.index(self.bulan_var.get())
            tahun = int(self.tahun_var.get())
        except Exception:
            messagebox.showerror("Export Excel", "Bulan atau tahun tidak valid.", parent=self)
            return

        periode_id = get_or_create_periode(bulan, tahun)
        default_name = f"Penugasan_{NAMA_BULAN[bulan]}_{tahun}.xlsx"
        path = filedialog.asksaveasfilename(
            parent=self,
            title="Simpan Export Penugasan",
            defaultextension=".xlsx",
            initialfile=default_name,
            filetypes=[("Excel Workbook", "*.xlsx")]
        )
        if not path:
            return

        try:
            wb = Workbook()
            conn = get_connection()
            try:
                self._export_calendar_sheet(wb, conn, bulan, tahun, periode_id)
                self._export_data_pegawai_sheet(wb, conn)
                bagian_rows = conn.execute(
                    """
                    SELECT id, kode, COALESCE(nama_bagian, ''), COALESCE(kegiatan, ''),
                           COALESCE(lokasi, ''), COALESCE(dana, 0), penanggung_jawab_id
                    FROM bagian
                    WHERE periode_id = ?
                    ORDER BY id
                    """,
                    (periode_id,)
                ).fetchall()
                for bagian in bagian_rows:
                    self._export_kegiatan_sheet(wb, conn, bagian, bulan, tahun)
            finally:
                conn.close()

            wb.save(path)
            wb.close()
            messagebox.showinfo(
                "Export Berhasil",
                f"File penugasan berhasil dibuat.\n\n{path}\n\n"
                f"Isi: Kalender Keseluruhan + Data Pegawai + {len(bagian_rows)} sheet kegiatan.",
                parent=self
            )
        except Exception as exc:
            messagebox.showerror(
                "Export Gagal",
                f"Gagal membuat file Excel.\n\n{exc}",
                parent=self
            )

    # =====================================================
    # KELOLA / BERSIHKAN DATA PENUGASAN
    # =====================================================

    def buka_kelola_data(self):
        """Jendela untuk menghapus histori penugasan per bulan tanpa menghapus master pegawai."""
        if getattr(self, "kelola_data_window", None) is not None:
            try:
                if self.kelola_data_window.winfo_exists():
                    self.kelola_data_window.lift()
                    self.kelola_data_window.focus_force()
                    self.refresh_kelola_data()
                    return
            except tk.TclError:
                pass

        win = tk.Toplevel(self)
        self.kelola_data_window = win
        win.title("Kelola / Bersihkan Data Penugasan")
        win.geometry("980x650")
        win.minsize(850, 560)
        win.transient(self)
        win.protocol("WM_DELETE_WINDOW", lambda: self.tutup_kelola_data(win))

        outer = tk.Frame(win, bg="#F4F6F8")
        outer.pack(fill="both", expand=True)
        tk.Label(outer, text="KELOLA / BERSIHKAN DATA PENUGASAN", bg="#F4F6F8", font=("Segoe UI", 18, "bold")).pack(anchor="w", padx=25, pady=(20, 5))
        tk.Label(
            outer,
            text=(
                "Data penugasan disimpan per bulan. Pilih satu bulan yang sudah tidak diperlukan "
                "untuk menghapus seluruh kegiatan, keanggotaan, jadwal, dan hari liburnya.\n"
                "Master Data Pegawai TIDAK ikut dihapus. Data yang sudah dihapus tidak dapat dipulihkan dari aplikasi."
            ),
            bg="#F4F6F8", fg="#555555", justify="left", font=("Segoe UI", 10)
        ).pack(anchor="w", padx=25, pady=(0, 15))

        table_frame = tk.Frame(outer, bg="white")
        table_frame.pack(fill="both", expand=True, padx=25, pady=5)
        self.kelola_data_tree = ttk.Treeview(
            table_frame, columns=("periode_id", "periode", "kegiatan", "anggota", "jadwal"),
            show="headings", selectmode="browse"
        )
        heads = {"periode_id":"ID", "periode":"Periode", "kegiatan":"Kegiatan", "anggota":"Keanggotaan", "jadwal":"Turun"}
        widths = {"periode_id":60, "periode":180, "kegiatan":130, "anggota":150, "jadwal":120}
        for c in heads:
            self.kelola_data_tree.heading(c, text=heads[c])
            self.kelola_data_tree.column(c, width=widths[c], anchor="center")
        self.kelola_data_tree.column("periode_id", width=60, stretch=False)
        self.kelola_data_tree.pack(side="left", fill="both", expand=True)
        sb = ttk.Scrollbar(table_frame, orient="vertical", command=self.kelola_data_tree.yview)
        sb.pack(side="right", fill="y")
        self.kelola_data_tree.configure(yscrollcommand=sb.set)
        self.kelola_data_tree.bind("<Delete>", lambda event: self.bersihkan_periode_terpilih())

        btn_frame = tk.Frame(outer, bg="#F4F6F8")
        btn_frame.pack(fill="x", padx=25, pady=(12, 20))
        ttk.Button(btn_frame, text="↻ Refresh", command=self.refresh_kelola_data).pack(side="left")
        ttk.Button(btn_frame, text="🗑 Bersihkan Bulan Terpilih", command=self.bersihkan_periode_terpilih).pack(side="left", padx=8)
        ttk.Button(btn_frame, text="⚠ Bersihkan SEMUA Histori", command=self.bersihkan_semua_histori).pack(side="right")
        self.refresh_kelola_data()

    def tutup_kelola_data(self, win=None):
        try:
            target = win or getattr(self, "kelola_data_window", None)
            if target is not None and target.winfo_exists():
                target.destroy()
        finally:
            self.kelola_data_window = None
            if hasattr(self, "kelola_data_tree"):
                self.kelola_data_tree = None

    def refresh_kelola_data(self):
        tree = getattr(self, "kelola_data_tree", None)
        if tree is None:
            return
        try:
            for item in tree.get_children():
                tree.delete(item)
        except tk.TclError:
            return
        conn = get_connection()
        try:
            rows = conn.execute("SELECT id, bulan, tahun FROM periode ORDER BY tahun DESC, bulan DESC").fetchall()
            for periode_id, bulan, tahun in rows:
                ringkas = get_ringkasan_periode(periode_id)
                tree.insert(
                    "", "end", iid=str(periode_id),
                    values=(periode_id, f"{NAMA_BULAN[int(bulan)]} {int(tahun)}", ringkas["kegiatan"], ringkas["anggota"], ringkas["jadwal"])
                )
        finally:
            conn.close()

    def bersihkan_periode_terpilih(self):
        tree = getattr(self, "kelola_data_tree", None)
        if tree is None:
            return
        pilihan = tree.selection()
        if not pilihan:
            messagebox.showinfo("Bersihkan Data", "Pilih bulan yang ingin dibersihkan terlebih dahulu.", parent=self.kelola_data_window)
            return
        periode_id = int(pilihan[0])
        ringkas = get_ringkasan_periode(periode_id)
        if not ringkas["periode"]:
            return
        bulan, tahun = int(ringkas["bulan"]), int(ringkas["tahun"])
        jawab = messagebox.askyesno(
            "Konfirmasi Bersihkan Bulan",
            f"Bersihkan seluruh data penugasan {NAMA_BULAN[bulan]} {tahun}?\n\n"
            f"• {ringkas['kegiatan']} kegiatan\n• {ringkas['anggota']} keanggotaan\n• {ringkas['jadwal']} kali turun/jadwal\n\n"
            "Yang dihapus: kegiatan, keanggotaan, jadwal, dan hari libur bulan tersebut.\n"
            "Yang tetap: Data Pegawai.\n\n"
            "Pastikan laporan Excel sudah disimpan jika masih diperlukan. Lanjutkan?",
            parent=self.kelola_data_window
        )
        if not jawab:
            return
        try:
            hapus_data_penugasan_periode(periode_id)
        except Exception as exc:
            messagebox.showerror("Gagal Membersihkan", f"Data gagal dibersihkan.\n\n{exc}", parent=self.kelola_data_window)
            return

        if periode_id == self.periode_id:
            self.periode_id = get_or_create_periode(self.bulan, self.tahun)
            self.bagian_aktif = load_bagian_codes(self.periode_id)[0]
            self.refresh_bagian_buttons()
            self.load_bagian_data()
            self.load_kalender()
            self.refresh_selection_table()
            self.refresh_pj_options()
            self.refresh_data_kegiatan()
            self.refresh_data_pegawai()
            self.load_kalender_keseluruhan()
        self.refresh_kelola_data()
        messagebox.showinfo("Berhasil", f"Data penugasan {NAMA_BULAN[bulan]} {tahun} sudah dibersihkan.\nMaster Data Pegawai tetap aman.", parent=self.kelola_data_window)

    def bersihkan_semua_histori(self):
        conn = get_connection()
        try:
            count = conn.execute("SELECT COUNT(*) FROM periode").fetchone()[0]
        finally:
            conn.close()
        if not count:
            messagebox.showinfo("Bersihkan Data", "Belum ada histori penugasan yang tersimpan.", parent=self.kelola_data_window)
            return
        if not messagebox.askyesno(
            "Bersihkan Semua Histori",
            "Semua data penugasan dari SEMUA BULAN akan dihapus.\n\n"
            "Yang dihapus: kegiatan, keanggotaan, jadwal, dan hari libur.\n"
            "Master Data Pegawai TIDAK dihapus.\n\n"
            "Pastikan semua laporan yang diperlukan sudah diekspor. Lanjutkan?",
            parent=self.kelola_data_window
        ):
            return
        konfirmasi = simpledialog.askstring("Konfirmasi Terakhir", "Ketik HAPUS untuk melanjutkan:", parent=self.kelola_data_window)
        if konfirmasi != "HAPUS":
            return
        try:
            hapus_semua_data_penugasan()
        except Exception as exc:
            messagebox.showerror("Gagal Membersihkan", f"Data gagal dibersihkan.\n\n{exc}", parent=self.kelola_data_window)
            return
        self.periode_id = get_or_create_periode(self.bulan, self.tahun)
        self.bagian_aktif = load_bagian_codes(self.periode_id)[0]
        self.refresh_bagian_buttons()
        self.load_bagian_data()
        self.load_kalender()
        self.refresh_selection_table()
        self.refresh_pj_options()
        self.refresh_data_kegiatan()
        self.refresh_data_pegawai()
        self.load_kalender_keseluruhan()
        self.refresh_kelola_data()
        messagebox.showinfo("Berhasil", "Semua histori penugasan sudah dibersihkan.\nMaster Data Pegawai tetap aman.", parent=self.kelola_data_window)

    def clear_search_data_pegawai(self):
        self.search_data_pegawai_var.set("")
        self.refresh_data_pegawai()

    def change_period(self, event=None):

        try:

            bulan = NAMA_BULAN.index(
                self.bulan_var.get()
            )

            tahun = int(
                self.tahun_var.get()
            )

        except:

            messagebox.showerror(
                "Periode",
                "Bulan atau tahun tidak valid."
            )

            return

        self.bulan = bulan
        self.tahun = tahun

        self.periode_id = get_or_create_periode(
            bulan,
            tahun
        )

        kode_list = load_bagian_codes(self.periode_id)
        if self.bagian_aktif not in kode_list:
            self.bagian_aktif = kode_list[0]

        self.refresh_bagian_buttons()

        self.header_period.config(
            text=f"{NAMA_BULAN[bulan]} {tahun}"
        )

        self.load_bagian_data()
        self.load_kalender()
        self.refresh_data_pegawai()

    def change_bagian(self, kode):
        self.bagian_aktif = kode
        self.load_bagian_data()
        self.load_kalender()
        self.refresh_data_kegiatan()
        self.update_bagian_button_state()

    def get_deskripsi(self):
        """Mengambil deskripsi dari kotak teks multi-baris."""
        if hasattr(self, "kegiatan_text"):
            return self.kegiatan_text.get("1.0", "end-1c").strip()
        return self.kegiatan_var.get().strip() if hasattr(self, "kegiatan_var") else ""

    def set_deskripsi(self, value):
        """Mengisi deskripsi ke kotak teks dan menjaga state tetap sinkron."""
        value = value or ""
        if hasattr(self, "kegiatan_text"):
            self.kegiatan_text.delete("1.0", "end")
            self.kegiatan_text.insert("1.0", value)
        if hasattr(self, "kegiatan_var"):
            self.kegiatan_var.set(value)

    def load_bagian_data(self):

        data = load_bagian(
            self.periode_id,
            self.bagian_aktif
        )

        if not data:
            return

        (
            bagian_id,
            kode,
            kegiatan,
            lokasi,
            dana,
            penanggung_jawab_id
        ) = data

        self.bagian_id = bagian_id

        conn = get_connection()
        nama_row = conn.execute(
            "SELECT COALESCE(nama_bagian, '') FROM bagian WHERE id = ?",
            (bagian_id,)
        ).fetchone()
        conn.close()

        nama_bagian = nama_row[0] if nama_row else ""
        self.nama_bagian_var.set(
            nama_bagian or f"Kegiatan {kode}"
        )

        deskripsi = kegiatan or ""
        # Sinkronkan state StringVar dan widget Text.
        self.kegiatan_var.set(deskripsi)
        self.set_deskripsi(deskripsi)

        self.lokasi_var.set(
            lokasi or ""
        )

        self.dana_var.set(
            str(dana or "")
        )

        self.penanggung_jawab_id = penanggung_jawab_id
        self.refresh_pj_options()
        self.set_pj_value(penanggung_jawab_id)

        anggota = load_anggota(bagian_id)

        self.selected_ids = set(
            pegawai_id
            for pegawai_id, nama in anggota
        )
    

        self.refresh_selection_table()
        self.refresh_data_kegiatan()

        nama_header = self.nama_bagian_var.get().strip() if hasattr(self, "nama_bagian_var") else f"Kegiatan {self.bagian_aktif}"
        self.header_period.config(
            text=f"{NAMA_BULAN[self.bulan]} {self.tahun}"
            f"  •  {nama_header}"
        )

        self.update_bagian_button_state()
        self.update_notebook_title()

    def refresh_pj_options(self):
        if not hasattr(self, "combo_pj"):
            return

        self.pj_options = {}
        values = []
        for p in self.pegawai:
            pegawai_id = p[0]
            nama = p[1] or "Tanpa Nama"
            nip = p[2] or "-"
            display = f"{nama} — {nip}"
            self.pj_options[display] = pegawai_id
            values.append(display)

        self.combo_pj["values"] = values

    def set_pj_value(self, pegawai_id):
        if not hasattr(self, "combo_pj"):
            return
        if not pegawai_id:
            self.pj_var.set("")
            return
        for display, pid in self.pj_options.items():
            if pid == pegawai_id:
                self.pj_var.set(display)
                return
        self.pj_var.set("")

    def get_selected_pj_id(self):
        return self.pj_options.get(self.pj_var.get())



# =========================================================
# RUN
# =========================================================

if __name__ == "__main__":

    app = App()

    app.mainloop()